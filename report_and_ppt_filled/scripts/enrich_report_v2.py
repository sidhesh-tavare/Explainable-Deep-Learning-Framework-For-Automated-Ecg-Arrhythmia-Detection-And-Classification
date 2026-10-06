"""Make a v2 copy of the filled mid-sem report with more figures, tables and metrics from the
V1/V2 notebooks (numbers from notebooks/V*/V*_analysis.md and the notebook CSVs).

Usage: python enrich_report_v2.py <filled.docx> <out_v2.docx>
The input file is never modified. Figures/tables are renumbered and in-text references updated.
"""
import re
import shutil
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

SRC, OUT = Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve()
ROOT = SRC.parent.parent  # FYP/
V1F = ROOT / "notebooks" / "V1" / "all_my_outputs" / "figures"
V2F = ROOT / "notebooks" / "V2" / "ECG_Arrhythmia_Complete_Results" / "figures"
FIG = SRC.parent / "figures"
assert SRC.resolve() != OUT.resolve()
shutil.copyfile(SRC, OUT)

# keep a copy of every figure used next to the report, like the first fill
new_figs = {
    "V1_class_distribution.png": V1F / "01_class_distribution.png",
    "V1_example_beats.png": V1F / "02_example_beats.png",
    "V1_training_curves.png": V1F / "03_training_curves.png",
    "V1_per_class_scores.png": V1F / "05_per_class_scores.png",
    "V2_01_augmentation_examples.png": V2F / "01_augmentation_examples.png",
    "V2_02_validation_curves.png": V2F / "02_validation_curves.png",
    "V2_03_experiment_comparison.png": V2F / "03_experiment_comparison.png",
    "V2_09_final_per_class_metrics.png": V2F / "09_final_per_class_metrics.png",
}
for name, src in new_figs.items():
    if not (FIG / name).exists():
        shutil.copyfile(src, FIG / name)

d = docx.Document(str(OUT))
TNR = "Times New Roman"
OLD_ELEMS = set(p._element for p in d.paragraphs)


# ---------- helpers (same style as fill_report.py) ----------------------------
def find(prefix):
    for p in d.paragraphs:
        if p.text.startswith(prefix) or prefix in p.text:
            return p
    raise KeyError(prefix)


def _run(p, text, bold=False, size=12):
    r = p.add_run(text)
    r.font.name = TNR
    r._element.rPr.rFonts.set(qn("w:hAnsi"), TNR)
    r._element.rPr.rFonts.set(qn("w:cs"), TNR)
    r.font.size = Pt(size)
    r.bold = bold or None
    return r


class Writer:
    def __init__(self, anchor):
        self.cur = anchor._element if hasattr(anchor, "_element") else anchor

    def _p(self):
        new = OxmlElement("w:p")
        self.cur.addnext(new)
        self.cur = new
        return Paragraph(new, d._body)

    def para(self, text, align="both"):
        p = self._p()
        p.alignment = {"both": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
        _run(p, text)
        return p

    def caption(self, text):
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        _run(p, text, bold=True)
        return p

    def figure(self, path, caption, width=6.0):
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Inches(width))
        self.caption(caption).paragraph_format.space_before = Pt(0)

    def table(self, caption, rows, widths, bold_rows=()):
        self.caption(caption).paragraph_format.keep_with_next = True
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = d.tables[0].style
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.width = Inches(widths[j])
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_after = Pt(0)
                cp.paragraph_format.keep_with_next = i < len(rows) - 1  # keep the table on one page
                cp.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                _run(cp, str(val), bold=(i == 0 or i in bold_rows), size=10)
        self.cur.addnext(t._tbl)
        self.cur = t._tbl
        self.para("")


def after_table(caption_prefix):
    """Spacer paragraph that follows the table under this caption."""
    tbl = find(caption_prefix)._element.getnext()
    assert tbl.tag == qn("w:tbl"), caption_prefix
    return tbl.getnext()


def after_figure(caption_prefix):
    return find(caption_prefix)._element


# ---------- 1. dataset: splits table + data figures -----------------------------
w = Writer(after_table("Table 2:"))
w.para("The split used in all preliminary experiments is summarised in Table {tab:split}. After the stratified "
       "split, the training portion contains 57,977 N, 1,778 S, 4,630 V, 513 F and 5,145 Q beats. Figure "
       "{fig:classdist} shows the class counts and Figure {fig:beats} one example beat of each class. Every beat "
       "starts at its R peak, contains a second large peak later in the window (most likely the next R peak) and "
       "is then zero-padded to 187 samples. The notebooks perform no filtering or segmentation of their own; this "
       "was done by the dataset authors [35].")
w.table("Table {tab:split}: Data split used in the preliminary experiments", [
    ["Set", "Beats", "Input shape", "Source", "Used for"],
    ["Training", "70,043", "(187, 1)", "80% of training file (stratified, seed 42)", "Fitting the weights"],
    ["Validation", "17,511", "(187, 1)", "20% of training file", "Epoch, model and threshold selection"],
    ["Test", "21,892", "(187, 1)", "Provided test file", "Final evaluation only"],
], [1.0, 0.8, 0.9, 2.2, 1.6])
w.figure(V1F / "01_class_distribution.png", "Figure {fig:classdist}: Class distribution of the training file (N ≈ 83%, F < 1%)", width=4.6)
w.figure(V1F / "02_example_beats.png", "Figure {fig:beats}: One example beat of each class (187 samples, "
         "125 Hz, zero-padded)", width=6.3)

# ---------- 2. V1: architecture + settings, curves, per-class bars -------------
w = Writer(find("V1: Baseline 1D CNN"))
w.para("Table {tab:arch} lists the layers of the baseline network and Table {tab:v1set} its training settings.")
w.table("Table {tab:arch}: Architecture of the baseline 1D CNN", [
    ["Layer", "Output shape", "Parameters"],
    ["Input (one beat)", "(187, 1)", "0"],
    ["Conv1D 32, k = 7 → BN → ReLU → MaxPool 2", "(93, 32)", "256 + 128"],
    ["Conv1D 64, k = 5 → BN → ReLU → MaxPool 2", "(46, 64)", "10,304 + 256"],
    ["Conv1D 128, k = 3 → BN → ReLU (Grad-CAM layer) → MaxPool 2", "(23, 128)", "24,704 + 512"],
    ["Global average pooling", "(128)", "0"],
    ["Dense 64, ReLU → Dropout 0.3", "(64)", "8,256"],
    ["Dense 5, softmax", "(5)", "325"],
    ["Total (trainable / non-trainable)", "", "44,741 (44,293 / 448)"],
], [3.4, 1.2, 1.6], bold_rows=(8,))
w.table("Table {tab:v1set}: Training settings of the V1 baseline", [
    ["Setting", "Value"],
    ["Optimiser / learning rate", "Adam, 10⁻³"],
    ["Loss", "Sparse categorical cross-entropy with balanced class weights"],
    ["Class weights (N, S, V, F, Q)", "0.24, 7.88, 3.03, 27.31, 2.72 (F:N ≈ 113:1)"],
    ["Batch size / maximum epochs", "128 / 30"],
    ["Early stopping", "Validation loss, patience 6, best weights restored"],
    ["Learning-rate schedule", "ReduceLROnPlateau on validation loss, factor 0.5, patience 3, min 10⁻⁵"],
    ["Augmentation", "None"],
    ["Random seed", "42 (Python, NumPy, TensorFlow)"],
], [2.2, 4.0])

w = Writer(find("The baseline reached 94.41% accuracy but only 76.29% macro F1"))
w.para("Figure {fig:v1curves} shows the training curves. Training loss fell smoothly from 0.706 to 0.175 and "
       "training accuracy was still rising (90.93% at epoch 17), while validation loss ranged from 0.349 "
       "(epoch 11, the restored epoch) to 4.29 (epoch 2). This gap between training and inference behaviour, "
       "rather than overfitting, is what V2 set out to fix. Figure {fig:v1scores} shows the per-class scores: "
       "the F and S precision bars are clearly below the recall bars, the typical sign of over-weighting the "
       "rare classes. The weighted F1 was 94.98% and macro precision / recall were 73.22% / 86.07%.")
w.figure(V1F / "03_training_curves.png", "Figure {fig:v1curves}: Training and validation accuracy and loss of the "
         "V1 baseline (17 epochs; best validation loss at epoch 11)", width=6.3)
w.figure(V1F / "05_per_class_scores.png", "Figure {fig:v1scores}: Per-class precision, recall and F1 of the V1 "
         "baseline on the test set", width=4.8)

# ---------- 3. V2: weights, configs, augmentation, training runs ---------------
w = Writer(find("Each V2 experiment keeps the same data split and seed as V1"))
w.para("Table {tab:weights} compares the two class-weighting schemes and Table {tab:cfg} lists the configuration "
       "of every experiment. All experiments use Adam, sparse categorical cross-entropy (with per-sample "
       "weights), batch size 128, at most 30 epochs and a custom training loop: after each epoch the validation "
       "macro F1 is computed, the best weights are stored, the learning rate is halved after every 3 epochs "
       "without improvement (minimum 10⁻⁵) and training stops after 8.")
w.table("Table {tab:weights}: Class weights used in the experiments", [
    ["Class", "Training beats", "Balanced weight (B0, E1, E2)", "Square-root weight (E3–E5)"],
    ["N", "57,977", "0.24", "0.49"],
    ["S", "1,778", "7.88", "2.81"],
    ["V", "4,630", "3.03", "1.74"],
    ["F", "513", "27.31", "5.23"],
    ["Q", "5,145", "2.72", "1.65"],
    ["F : N ratio", "", "113 : 1", "10.6 : 1"],
], [1.1, 1.3, 1.9, 1.9], bold_rows=(6,))
w.table("Table {tab:cfg}: Configuration of the V2 experiments (each keeps all earlier changes)", [
    ["ID", "Network", "Params", "LR", "Clip norm", "BN momentum", "Class weights", "Augment."],
    ["E1", "CNN", "44,741", "1×10⁻³", "–", "0.99", "Balanced", "No"],
    ["E2", "CNN", "44,741", "5×10⁻⁴", "1.0", "0.9", "Balanced", "No"],
    ["E3", "CNN", "44,741", "5×10⁻⁴", "1.0", "0.9", "Square-root", "No"],
    ["E4", "CNN", "44,741", "5×10⁻⁴", "1.0", "0.9", "Square-root", "Yes"],
    ["E5", "Residual CNN", "185,477", "5×10⁻⁴", "1.0", "0.9", "Square-root", "Yes"],
], [0.45, 1.05, 0.75, 0.7, 0.65, 0.85, 1.0, 0.75])
w.para("The residual network (E5) uses a Conv1D 32 (k = 7) stem followed by three residual blocks with 32, 64 and "
       "128 filters; each block has two Conv1D (k = 5) + BN layers added to a shortcut (a 1×1 convolution when "
       "the channel count changes), then ReLU and max-pooling, and the same pooling and dense head as the "
       "baseline. The augmentation of E4 and E5 is applied afresh to every training beat in every epoch: "
       "amplitude scaling by U(0.9, 1.1), Gaussian noise (σ = 0.01) added only to the non-padded part, and a "
       "time shift of −3 to +3 samples (±24 ms). Figure {fig:aug} shows that the augmented copies stay close "
       "to the original beat.")
w.figure(V2F / "01_augmentation_examples.png", "Figure {fig:aug}: Original and three augmented copies of an N, "
         "S and F beat", width=6.3)
w.table("Table {tab:runs}: Training runs of the V2 experiments", [
    ["ID", "Best epoch / epochs run", "Stopped early", "Best val. macro F1 (%)", "Δ test macro F1", "Time (s)"],
    ["E1", "16 / 24", "Yes", "85.43", "+7.67 (vs B0)", "70"],
    ["E2", "28 / 30", "No (cap)", "79.35", "−5.66", "80"],
    ["E3", "23 / 30", "No (cap)", "88.65", "+9.28", "80"],
    ["E4", "25 / 30", "No (cap)", "88.01", "−1.38", "95"],
    ["E5", "18 / 26", "Yes", "93.19", "+4.57", "158"],
], [0.5, 1.3, 1.0, 1.4, 1.2, 0.8])
w.figure(V2F / "02_validation_curves.png", "Figure {fig:valcurves}: Validation macro F1 and validation accuracy "
         "per epoch for E1–E5", width=6.3)
w.para("Table {tab:runs} summarises each training run. Figure {fig:valcurves} shows that E5 trains highest and most smoothly (validation macro F1 0.78 after "
       "the first epoch, validation accuracy between 0.94 and 0.99), E1 is the most jagged, and E2 is the "
       "lowest; E2, E3 and E4 all reached the 30-epoch cap, so they may be slightly under-trained. All five "
       "experiments together took 8.1 minutes on a Tesla T4 GPU.")

w = Writer(after_table("Table 4:"))
w.para("Figure {fig:expcmp} plots test macro F1 together with S and F recall for each step, Figure "
       "{fig:overall} compares test accuracy and macro F1, and Table {tab:f1all} gives the F1 of every class.")
w.figure(V2F / "03_experiment_comparison.png", "Figure {fig:expcmp}: Test macro F1, S recall and F recall "
         "for each step", width=5.6)
w.figure(V2F / "07_overall_model_comparison.png", "Figure {fig:overall}: Test accuracy and macro F1 from the "
         "baseline (B0) to the final model", width=5.6)
w.table("Table {tab:f1all}: Per-class test F1 (%) of every model", [
    ["Model", "Acc.", "Macro F1", "N", "S", "V", "F", "Q"],
    ["B0", "94.41", "76.29", "97.00", "60.90", "86.60", "40.00", "97.00"],
    ["E1", "96.28", "83.96", "97.92", "63.48", "92.51", "68.12", "97.77"],
    ["E2", "93.58", "78.30", "96.27", "49.13", "90.77", "58.49", "96.84"],
    ["E3", "97.39", "87.58", "98.58", "69.84", "94.19", "76.97", "98.29"],
    ["E4", "97.02", "86.20", "98.40", "66.16", "91.93", "76.25", "98.29"],
    ["E5", "98.29", "90.77", "99.13", "81.22", "95.70", "78.53", "99.28"],
    ["Final", "98.41", "91.10", "99.18", "82.14", "95.87", "79.23", "99.06"],
], [0.7, 0.7, 0.85, 0.7, 0.7, 0.7, 0.7, 0.7], bold_rows=(7,))

# ---------- 4. final model: bars, selection details, baseline-vs-final ----------
w = Writer(after_table("Table 5:"))
w.para("The final model is E5, which had the highest validation macro F1 (93.19%). The probability scale "
       "factors tuned on the validation set were N = 3.0 and 1.0 for every other class, so the model must be "
       "more confident before predicting any non-N class. This raised test accuracy from 98.29% to 98.41% and "
       "macro F1 from 90.77% to 91.10%. The weighted-average F1 is 98.37%. Figure {fig:finalbars} shows the "
       "per-class scores.")
w.figure(V2F / "09_final_per_class_metrics.png", "Figure {fig:finalbars}: Per-class precision, recall and F1 "
         "of the final model on the test set", width=5.2)

w = Writer(find("Compared with the baseline, accuracy rose from 94.41% to 98.41%"))
w.para("Table {tab:b0final} summarises the change from the baseline to the final model.")
w.table("Table {tab:b0final}: Baseline (V1) and final model compared on the same 21,892-beat test set", [
    ["Metric", "Baseline (B0)", "Final model", "Change"],
    ["Accuracy (%)", "94.41", "98.41", "+4.00"],
    ["Macro F1 (%)", "76.29", "91.10", "+14.81"],
    ["Weighted F1 (%)", "94.98", "98.37", "+3.39"],
    ["Precision S / Recall S (%)", "55.86 / 66.91", "87.58 / 77.34", "+31.72 / +10.43"],
    ["Precision F / Recall F (%)", "26.02 / 86.42", "82.12 / 76.54", "+56.10 / −9.88"],
    ["Beats wrongly predicted as F", "398", "27", "−371"],
    ["Beats wrongly predicted as S", "294", "61", "−233"],
    ["S beats predicted as N", "163 (29%)", "120 (22%)", "−43"],
    ["Parameters", "44,741", "185,477", "≈ 4×"],
], [2.4, 1.3, 1.3, 1.3])

# ---------- 5. Grad-CAM: region description table -------------------------------
w = Writer(find("Both models concentrate their attention on a few narrow regions"))
w.para("Table {tab:gradcam} describes the heatmaps of both models class by class.")
w.table("Table {tab:gradcam}: Where the Grad-CAM heat is concentrated (read from the heatmap figures; "
        "all beats predicted with p = 1.00)", [
    ["Class", "V1 baseline CNN", "Final model (residual CNN)"],
    ["N", "Window start (0–0.1 s) and 1.15–1.3 s around the second peak / start of padding",
     "0.12–0.2 s and 0.4–0.5 s, before the second R peak (P-wave-like bump); second peak itself low"],
    ["S", "0.35–0.45 s over the second (sharp) peak and the drop into padding",
     "Window start (0–0.2 s) and 0.75–0.9 s just after the second peak"],
    ["V", "Window start, broad 0.3–0.9 s region, strong band at 1.0–1.1 s where the signal ends",
     "Window start, 0.27–0.35 s, and 1.15–1.3 s around the second peak / end of signal"],
    ["F", "Window start and two narrow bands either side of the second peak (~0.4 and ~0.53 s)",
     "One strong broad band at 0.3–0.5 s centred on the second peak"],
    ["Q", "Window start and a strong band at 0.75–0.85 s over the second peak",
     "~0.1 s (deep trough) and 0.75–0.95 s around the second peak"],
], [0.6, 2.8, 2.8])
w.para("Grad-CAM is computed from a layer with 46 time steps, so each step covers about 4 samples (≈ 32 ms) "
       "before interpolation; strong heat at the very first samples may partly be an edge effect of the "
       "'same' padding. Because the length of the non-zero part of each window depends on heart rate, heat at "
       "the start of the padding may also reflect beat timing. Both points will be checked with the "
       "region-share and deletion measures.")

# ---------- renumber figures and tables, update references ---------------------
CAP = re.compile(r"^(Figure|Table) (\d+|\{(?:fig|tab):\w+\}):")
counter = {"Figure": 0, "Table": 0}
old_map = {"Figure": {}, "Table": {}}
tok_map = {}
for p in d.paragraphs:
    m = CAP.match(p.text)
    if not m:
        continue
    kind, key = m.groups()
    counter[kind] += 1
    n = counter[kind]
    if key.startswith("{"):
        tok_map[key] = str(n)
    else:
        old_map[kind][key] = str(n)
print("figures:", counter["Figure"], "tables:", counter["Table"])
print("old->new:", old_map)

REF = re.compile(r"\b(Figures?|Tables?) (\d+)((?:,\s*\d+)*)((?: and | to )(\d+))?")


def fix_old(text):
    def sub(m):
        kind = "Figure" if m.group(1).startswith("F") else "Table"
        mp = old_map[kind]
        out = m.group(1) + " " + mp.get(m.group(2), m.group(2))
        if m.group(3):
            out += re.sub(r"\d+", lambda x: mp.get(x.group(0), x.group(0)), m.group(3))
        if m.group(4):
            out += m.group(4).replace(m.group(5), "") + mp.get(m.group(5), m.group(5))
        return out
    return REF.sub(sub, text)


def fix_tok(text):
    return re.sub(r"\{(?:fig|tab):\w+\}", lambda m: tok_map[m.group(0)], text)


def all_paragraphs():
    yield from d.paragraphs
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                yield from c.paragraphs


for p in all_paragraphs():
    is_old = p._element in OLD_ELEMS
    for r in p.runs:
        t = r.text
        if is_old:
            t = fix_old(t)  # old captions and old in-text references
        t = fix_tok(t)
        if t != r.text:
            r.text = t

d.save(str(OUT))
left = [p.text for p in all_paragraphs() if "{fig:" in p.text or "{tab:" in p.text or "TO FILL" in p.text]
print("saved", OUT.name, "| unresolved:", left)
