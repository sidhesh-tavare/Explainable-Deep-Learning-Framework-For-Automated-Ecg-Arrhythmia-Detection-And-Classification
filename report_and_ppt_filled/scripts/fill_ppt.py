"""Add the V1/V2 results to the mid-sem deck, matching the existing slide style."""
import copy
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

PPT = Path(sys.argv[1])
FIG = PPT.parent / "figures"
prs = Presentation(str(PPT))
DARK, GREY, NAVY, ORANGE = "202020", "595959", "1F3864", "ED7D31"
TABLE_STYLE = "{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}"


# ---------- helpers ----------------------------------------------------------
def is_bar(sh):
    return sh.width == prs.slide_width and (sh.top < Emu(200000) or sh.top > Emu(6600000))


def is_title(sh):
    return (not is_bar(sh) and sh.has_text_frame and sh.text_frame.text.strip() != ""
            and sh.top < Emu(500000) and sh.width > Emu(4000000))


def set_text(shape, text):
    """Replace a text box's text, keeping the first run's formatting."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    p0.runs[0].text = text


def set_lines(shape, lines):
    """Replace text with several paragraphs, each copying the first run's formatting."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    tmpl_p = copy.deepcopy(p0._p)
    for p in list(tf.paragraphs):
        p._p.getparent().remove(p._p)
    for line in lines:
        np_ = copy.deepcopy(tmpl_p)
        runs = np_.findall(qn("a:r"))
        for r in runs[1:]:
            np_.remove(r)
        runs[0].find(qn("a:t")).text = line
        tf._txBody.append(np_)


def textbox(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, autofit=True):
    """paras: list of (text, size, bold, color, italic) or list of lists of such runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    bp = tf._txBody.find(qn("a:bodyPr"))
    for k in ("lIns", "tIns", "rIns", "bIns"):
        bp.set(k, "0")
    if autofit and bp.find(qn("a:spAutoFit")) is None:
        etree.SubElement(bp, qn("a:spAutoFit"))
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = para if isinstance(para, list) else [para]
        for text, size, bold, color, italic in runs:
            r = p.add_run()
            r.text = text
            f = r.font
            f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, "Calibri"
            f.color.rgb = RGBColor.from_string(color)
    return tb


def heading(slide, x, y, w, text, size=15):
    return textbox(slide, x, y, w, 0.35, [(text, size, True, DARK, False)])


def body(slide, x, y, w, h, lines, size=12, color=DARK):
    paras = []
    for ln in lines:
        if isinstance(ln, list):
            paras.append([(t, size, b, color, False) for t, b in ln])
        else:
            paras.append((ln, size, False, color, False))
    return textbox(slide, x, y, w, h, paras)


def note(slide, x, y, w, text, size=11):
    return textbox(slide, x, y, w, 0.3, [(text, size, False, GREY, True)])


def table(slide, x, y, w, rows, col_w, row_h=0.27, size=11, hl_rows=()):
    gf = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                Inches(row_h * len(rows)))
    tbl = gf.table
    tbl._tbl.tblPr.find(qn("a:tableStyleId")).text = TABLE_STYLE
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.margin_left = c.margin_right = Inches(0.06)
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = RGBColor.from_string(NAVY)
            elif i in hl_rows:
                c.fill.fore_color.rgb = RGBColor.from_string("FCE4D6")
            else:
                c.fill.fore_color.rgb = RGBColor.from_string("F2F2F2" if i % 2 else "FFFFFF")
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 or (j == 1 and len(row) > 6) else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(val)
            f = r.font
            f.size, f.name = Pt(size), "Calibri"
            f.bold = i == 0 or j == 0 or i in hl_rows
            f.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else DARK)
    return gf


def picture(slide, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)


TEMPLATE = prs.slides[10]  # "Current Work" slide: bars + title to copy
BLANK = TEMPLATE.slide_layout


def new_slide(title):
    s = prs.slides.add_slide(BLANK)
    for sh in TEMPLATE.shapes:
        if is_bar(sh) or is_title(sh):
            s.shapes._spTree.append(copy.deepcopy(sh._element))
    for sh in s.shapes:
        if is_title(sh):
            set_text(sh, title)
    return s


def clear_content(slide):
    for sh in list(slide.shapes):
        if not (is_bar(sh) or is_title(sh)):
            sh._element.getparent().remove(sh._element)


# ---------- slide 11: work done overview --------------------------------------
s11 = prs.slides[10]
clear_content(s11)
for sh in s11.shapes:
    if is_title(sh):
        set_text(sh, "Work Done Till Mid-Semester")

heading(s11, 0.55, 1.3, 4.1, "Completed")
body(s11, 0.55, 1.72, 4.1, 2.2, [
    "• Literature review of 37 papers; research gap identified",
    "• Framework, splits (incl. DS1 / DS2) and metrics defined",
    "• V1: baseline 1D CNN, class-weighted loss, Grad-CAM",
    "• V2: 5 step-by-step experiments + threshold tuning",
    [("• Final model: ", False), ("98.41% accuracy, 91.10% macro-F1", True)],
])
heading(s11, 5.0, 1.3, 4.4, "Development dataset [13]")
body(s11, 5.0, 1.72, 4.4, 2.2, [
    "Kaggle version of MIT-BIH, already segmented by the dataset authors: 187-sample beats "
    "(1.5 s at 125 Hz), scaled 0–1, five AAMI classes (N, S, V, F, Q).",
    "Used to build and debug the training, evaluation and Grad-CAM pipeline before our own "
    "MIT-BIH preprocessing.",
])
table(s11, 0.65, 4.2, 8.2, [
    ["File", "N", "S", "V", "F", "Q", "Total"],
    ["Train file", "72,471", "2,223", "5,788", "641", "6,431", "87,554"],
    ["Test file", "18,118", "556", "1,448", "162", "1,608", "21,892"],
], [1.6, 1.1, 1.1, 1.1, 1.1, 1.1, 1.1], row_h=0.32, size=12)
note(s11, 0.65, 5.3, 8.2,
     "Train file split 80/20 (stratified, seed 42) → 70,043 train / 17,511 validation beats; the test file is "
     "used only for final scores.")
body(s11, 0.65, 5.9, 8.2, 0.6, [
    [("Caveat: ", True), ("the split is by beat, so the same patients appear in training and test data "
                         "(intra-patient). These scores are optimistic; DS1 / DS2 comes next.", False)]])

# ---------- slide: V1 results --------------------------------------------------
s = new_slide("V1 Result: Baseline 1D CNN")
picture(s, FIG / "V1_confusion_matrix.png", 0.9, 1.15, w=8.2)
table(s, 0.55, 4.5, 4.3, [
    ["Class", "Precision", "Recall", "F1"],
    ["N", "0.981", "0.959", "0.970"],
    ["S", "0.559", "0.669", "0.609"],
    ["V", "0.871", "0.861", "0.866"],
    ["F", "0.260", "0.864", "0.400"],
    ["Q", "0.990", "0.950", "0.970"],
    ["Macro", "0.732", "0.861", "0.763"],
], [1.0, 1.1, 1.1, 1.1], row_h=0.27, size=11, hl_rows=(6,))
heading(s, 5.15, 4.45, 4.3, "Accuracy 94.41%  |  Macro-F1 76.29%")
body(s, 5.15, 4.9, 4.3, 2.0, [
    "• 538 beats predicted F, only 140 truly F (293 were N) → F precision 26%",
    "• 29% of S beats called N (163 / 556)",
    "• Validation accuracy swung 14.5% ↔ 95.1%; early stop at epoch 17 (best: 11)",
    "• Cause: balanced weights make 1 F beat count as 113 N beats",
], size=11)

# ---------- slide: V2 experiments ---------------------------------------------
s = new_slide("V2: Step-by-Step Improvements")
note(s, 0.55, 1.15, 8.9, "Same data, split and seed as V1. One change added per step; epoch, model and "
                          "thresholds chosen on validation only.", size=12)
table(s, 0.55, 1.6, 8.9, [
    ["ID", "Change added", "Acc. %", "Macro-F1 %", "S F1 %", "F F1 %", "Δ Macro-F1"],
    ["B0", "V1 baseline", "94.41", "76.29", "60.90", "40.00", "–"],
    ["E1", "Best epoch by val. macro-F1", "96.28", "83.96", "63.48", "68.12", "+7.67"],
    ["E2", "+ LR 5e-4, clipping, BN mom. 0.9", "93.58", "78.30", "49.13", "58.49", "−5.66"],
    ["E3", "+ square-root class weights", "97.39", "87.58", "69.84", "76.97", "+9.28"],
    ["E4", "+ light augmentation", "97.02", "86.20", "66.16", "76.25", "−1.38"],
    ["E5", "+ residual CNN [14] (185 k params)", "98.29", "90.77", "81.22", "78.53", "+4.57"],
    ["Final", "E5 + tuned thresholds (N × 3)", "98.41", "91.10", "82.14", "79.23", "+0.33"],
], [0.7, 3.0, 0.95, 1.15, 0.95, 0.95, 1.2], row_h=0.32, size=11, hl_rows=(7,))
heading(s, 0.55, 4.45, 4.2, "What helped")
body(s, 0.55, 4.85, 4.2, 1.8, [
    "• E1: picking the epoch on macro-F1, not val. loss",
    "• E3: softer weights (F:N 113:1 → 10.6:1) – biggest gain, F precision 44 → 76%",
    "• E5: residual network – smoothest training, S F1 66 → 81",
], size=11)
heading(s, 5.15, 4.45, 4.3, "What did not")
body(s, 5.15, 4.85, 4.3, 1.8, [
    "• E2 stability settings: −5.66, swings remained",
    "• E4 augmentation: −1.38",
    "• Threshold tuning: only +0.33, trades S / F recall for precision",
    "• Steps are cumulative, so E2 / E4 were kept in the final model",
], size=11)

# ---------- slide: final model --------------------------------------------------
s = new_slide("Final Model: Residual CNN + Tuned Thresholds")
picture(s, FIG / "V2_04_confusion_matrix_final.png", 0.9, 1.15, w=8.2)
table(s, 0.55, 4.5, 4.3, [
    ["Class", "Precision", "Recall", "F1"],
    ["N", "0.987", "0.996", "0.992"],
    ["S", "0.876", "0.773", "0.821"],
    ["V", "0.981", "0.938", "0.959"],
    ["F", "0.821", "0.765", "0.792"],
    ["Q", "0.998", "0.983", "0.991"],
    ["Macro", "0.933", "0.891", "0.911"],
], [1.0, 1.1, 1.1, 1.1], row_h=0.27, size=11, hl_rows=(6,))
table(s, 5.15, 4.5, 4.3, [
    ["Metric (%)", "V1 (B0)", "Final"],
    ["Accuracy", "94.41", "98.41"],
    ["Macro-F1", "76.29", "91.10"],
    ["Precision S", "55.86", "87.58"],
    ["Precision F", "26.02", "82.12"],
    ["Recall S", "66.91", "77.34"],
    ["Recall F", "86.42", "76.54"],
], [1.7, 1.3, 1.3], row_h=0.27, size=11, hl_rows=(2,))
note(s, 0.55, 6.5, 8.9, "Largest remaining error: S → N (120 of 556 S beats). S differs from N mainly in timing, "
                         "which motivates the RR features. F recall fell because V1's high recall came from "
                         "over-predicting F.", size=10.5)

# ---------- slide: Grad-CAM -----------------------------------------------------
s = new_slide("Grad-CAM Explanations")
heading(s, 0.55, 1.1, 8.9, "V1 baseline CNN", size=13)
picture(s, FIG / "V1_gradcam.png", 0.55, 1.4, w=8.9)
heading(s, 0.55, 3.3, 8.9, "Final model (residual CNN)", size=13)
picture(s, FIG / "V2_06_gradcam_final.png", 0.55, 3.6, w=8.9)
body(s, 0.55, 5.6, 8.9, 1.3, [
    "• Most confident correct test beat per class; red = regions that drove the prediction.",
    "• Heat concentrates at the window start and around the second R peak → the model likely uses the timing "
    "of the next beat (an implicit RR cue), not only beat shape. Final-model maps are broader.",
    "• Qualitative only so far → next: P / QRS / T region share and deletion test, plus SHAP.",
], size=11)

# ---------- slide: challenges --------------------------------------------------
s = new_slide("Challenges & Lessons Learned")
cards = [
    ("Class imbalance", "Balanced weights (F:N 113:1) made the model over-predict F: precision 26%. "
                        "Square-root weights (10.6:1) fixed most of it."),
    ("Unstable training", "Validation accuracy swung 14.5% ↔ 95.1%; early stopping on validation loss "
                          "stopped too early. Selecting on macro-F1 and a residual CNN helped."),
    ("Not every change helps", "Stability settings (E2) and augmentation (E4) lowered macro-F1. Steps were "
                               "cumulative, so the best combination is not yet known."),
    ("Optimistic setting", "Pre-segmented Kaggle beats, intra-patient split, 5 classes, single seed. Next: raw "
                           "MIT-BIH, 4 classes, RR features and DS1 / DS2."),
]
for k, (h, t) in enumerate(cards):
    x = 0.65 if k % 2 == 0 else 5.0
    y = 1.45 if k < 2 else 3.95
    heading(s, x, y, 4.0, h)
    body(s, x, y + 0.42, 4.0, 1.6, [t], size=12)

# ---------- move new slides after slide 11 -----------------------------------
ids = prs.slides._sldIdLst
new = list(ids)[-5:]
for el in new:
    ids.remove(el)
for i, el in enumerate(new):
    ids.insert(11 + i, el)

# ---------- conclusion -------------------------------------------------------
concl = prs.slides[16]
boxes = {sh.text_frame.text.split("\n")[0][:20]: sh for sh in concl.shapes if sh.has_text_frame}
for sh in concl.shapes:
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text
    if t.startswith("• Published accuracies"):
        set_lines(sh, [
            "• A working pipeline is in place: training, macro-F1-based evaluation and Grad-CAM.",
            "• Macro-F1 76.29% → 91.10% (accuracy 94.41% → 98.41%) via epoch selection, softer class weights "
            "and a residual CNN.",
            "• S → N remains the main error; results are intra-patient, so inter-patient scores will be lower.",
            "• Published 98–99% accuracies are mostly intra-patient; our framework adds region, deletion and "
            "noise-stability scores.",
        ])
    elif t.startswith("Baseline CNN →"):
        set_text(sh, "Raw MIT-BIH preprocessing (4 classes, RR features) → CNN vs CNN-LSTM + RR → DS1 / DS2 "
                     "evaluation → Grad-CAM / SHAP with region and deletion checks → NSTDB noise tests → "
                     "Streamlit demo.")

# ---------- references -------------------------------------------------------
refs = prs.slides[17]
for sh in refs.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("[1]"):
        tf = sh.text_frame
        last = tf.paragraphs[-1]._p
        for txt in [
            "[13] M. Kachuee, S. Fazeli and M. Sarrafzadeh, “ECG heartbeat classification: A deep transferable "
            "representation,” Proc. IEEE ICHI, pp. 443–444, 2018.",
            "[14] K. He, X. Zhang, S. Ren and J. Sun, “Deep residual learning for image recognition,” "
            "Proc. IEEE CVPR, pp. 770–778, 2016.",
        ]:
            np_ = copy.deepcopy(last)
            rs = np_.findall(qn("a:r"))
            for r in rs[1:]:
                np_.remove(r)
            rs[0].find(qn("a:t")).text = txt
            last.addnext(np_)
            last = np_

prs.save(str(PPT))
print("slides:", len(prs.slides))
