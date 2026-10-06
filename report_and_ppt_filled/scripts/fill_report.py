"""Fill the [TO FILL] parts of the mid-sem report with the V1/V2 notebook results."""
import copy
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

DOC = Path(sys.argv[1])
FIG = DOC.parent / "figures"
d = docx.Document(str(DOC))
body = d.element.body
TNR = "Times New Roman"


# ---------- helpers ----------------------------------------------------------
def find(prefix):
    for p in d.paragraphs:
        if prefix in p.text:
            return p
    raise KeyError(prefix)


def _run(p, text, bold=False, italic=False, size=12):
    r = p.add_run(text)
    r.font.name = TNR
    r._element.rPr.rFonts.set(qn("w:hAnsi"), TNR)
    r._element.rPr.rFonts.set(qn("w:cs"), TNR)
    r.font.size = Pt(size)
    r.bold = bold or None
    r.italic = italic or None
    return r


def _new_p_after(anchor_el):
    new = OxmlElement("w:p")
    anchor_el.addnext(new)
    return Paragraph(new, d._body)


class Writer:
    """Appends paragraphs/tables/figures one after another, starting after an anchor."""

    def __init__(self, anchor):
        self.cur = anchor._element if hasattr(anchor, "_element") else anchor

    def _p(self):
        p = _new_p_after(self.cur)
        self.cur = p._element
        return p

    def para(self, parts, align="both"):
        """parts: str or list of (text, bold) tuples."""
        p = self._p()
        p.alignment = {"both": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
        if isinstance(parts, str):
            parts = [(parts, False)]
        for t, b in parts:
            _run(p, t, bold=b)
        return p

    def sub(self, text):
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.keep_with_next = True
        _run(p, text, bold=True)
        return p

    def bullet(self, text, lead=None):
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.left_indent = Inches(0.3)
        pf.first_line_indent = Inches(-0.3)
        pf.tab_stops.add_tab_stop(Inches(0.3))
        _run(p, "•\t")
        if lead:
            _run(p, lead, bold=True)
        _run(p, text)
        return p

    def caption(self, text):
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        _run(p, text, bold=True)
        return p

    def figure(self, path, caption, width=6.0):
        p = self._p()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Inches(width))
        c = self.caption(caption)
        c.paragraph_format.space_before = Pt(0)
        return c

    def table(self, caption, rows, widths, bold_rows=()):
        self.caption(caption).paragraph_format.keep_with_next = True
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = d.tables[0].style  # same "TableGrid" style as Table 1
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.width = Inches(widths[j])
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_after = Pt(0)
                cp.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                _run(cp, str(val), bold=(i == 0 or i in bold_rows), size=10)
        self.cur.addnext(t._tbl)
        self.cur = t._tbl
        self.para("")  # spacer after the table
        return t


def replace_text(p, text, bold=False):
    """Replace a paragraph's runs with plain text, keeping its paragraph format."""
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    _run(p, text, bold=bold)


# ---------- title page ------------------------------------------------------
p = find("[TO FILL: NAME OF THE STUDENTS")
for r in list(p.runs):
    r._element.getparent().remove(r._element)
r = _run(p, "SIDHESH MAHESH TAVARE (2314037)", bold=True, size=14)
r.add_break()
_run(p, "NIKHIL KUMAR SINGH (2314055)", bold=True, size=14)

p = find("[TO FILL: SUPERVISOR NAME]")
replace_text(p, "DR. ARNAB NANDI", bold=True)
p.runs[0].font.size = Pt(14)

# ---------- abstract: add a progress paragraph ------------------------------
w = Writer(find("The expected outcomes are a patient-independent"))
w.para("Up to the mid-semester evaluation, a working training, evaluation and Grad-CAM pipeline has been built "
       "and tested on the pre-segmented MIT-BIH heartbeat dataset of Kachuee et al. [35] (five AAMI classes, "
       "intra-patient split). A baseline 1D convolutional neural network reached 94.41% accuracy but only 76.29% "
       "macro F1, because of very low precision on the rare Fusion (26.0%) and supraventricular (55.9%) classes. "
       "Five step-by-step changes were then tested; selecting the epoch on validation macro F1, using softer "
       "(square-root) class weights and a residual network gave the largest gains, and the final model reached "
       "98.41% accuracy and 91.10% macro F1 on the same 21,892-beat test set. Grad-CAM heatmaps were produced "
       "for every class. These results are intra-patient and are therefore treated as a development baseline "
       "for the patient-independent framework.")

# ---------- methodology -----------------------------------------------------
p = find("The proposed framework follows directly from the three objectives")
replace_text(p, "The proposed framework follows directly from the three objectives. It consists of the stages "
                "below and is summarised in Figure 1. Choices that were tested in the preliminary experiments are "
                "noted in each subsection; the experiments themselves are described in the next chapter.")

p = find("[TO FILL: insert block diagram figure]")
for r in list(p.runs):
    r._element.getparent().remove(r._element)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.keep_with_next = True
p.add_run().add_picture(str(FIG / "block_diagram.png"), width=Inches(6.0))
Writer(p).caption("Figure 1: Block diagram of the proposed framework")

p = find("An external database for cross-database testing")
for r in p.runs:
    if "[TO FILL" in r.text:
        r._element.getparent().remove(r._element)
_run(p, "The INCART database is the preferred candidate because, like MIT-BIH, it provides beat-level "
        "annotations; external testing is treated as a stretch goal and is otherwise listed as future work.")

p = find("A hybrid model combining a convolutional front end")
for r in p.runs:
    if "[TO FILL" in r.text:
        r._element.getparent().remove(r._element)
w = Writer(p)
w.bullet("a 1D CNN with three convolutional blocks (32, 64 and 128 filters, kernel sizes 7, 5 and 3; each "
         "followed by batch normalisation, ReLU and max-pooling), global average pooling, a 64-unit dense layer "
         "and a softmax output (about 44,700 parameters). This network has already been implemented and "
         "trained (next chapter).", lead="Baseline: ")
w.bullet("the same convolutional front end followed by an LSTM layer with 64 units; its last hidden state is "
         "concatenated with the previous and next RR intervals (each divided by the record's mean RR) before a "
         "64-unit dense layer and a softmax over N, S, V and F. A residual CNN, which gave the best preliminary "
         "results, will be kept as an additional comparison model.", lead="Proposed: ")
w.bullet("class-weighted cross-entropy on the training partition only. The preliminary experiments showed that "
         "fully balanced weights (F:N weight ratio 113:1) caused many false Fusion predictions, while "
         "square-root-balanced weights (10.6:1) gave the best results, so square-root weights are used.",
         lead="Class imbalance: ")
w.bullet("Adam optimiser, batch size 128, at most 30 epochs, with the best epoch chosen on validation macro F1 "
         "rather than validation loss (this alone raised test macro F1 from 76.29% to 83.96% in the preliminary "
         "experiments). The learning rate will be tuned on the validation set.", lead="Training: ")

p = find("At least two post-hoc XAI methods")
for r in p.runs:
    if "[TO FILL" in r.text:
        r._element.getparent().remove(r._element)
_run(p, "The two methods are Grad-CAM [37], already implemented for the preliminary models, and Gradient SHAP [38]. "
        "Three measures will be used:")
w = Writer(p)
w.bullet("the share of total absolute attribution that falls in fixed windows around the R peak: P wave "
         "[−250, −80] ms, QRS complex [−50, +50] ms and T wave [+100, +400] ms, averaged per class.",
         lead="Region share: ")
w.bullet("the 10% of samples with the highest attribution are set to zero and the drop in the predicted class "
         "probability is compared with the drop when 10% of samples are removed at random. A faithful map should "
         "give a much larger drop.", lead="Deletion test: ")
w.bullet("the Pearson correlation between the attribution map of a clean beat and of the same beat with noise "
         "added (used in the robustness study).", lead="Stability: ")

p = find("Classification metrics and explanation stability will be measured")
for r in p.runs:
    if "[TO FILL" in r.text:
        r._element.getparent().remove(r._element)
_run(p, "The muscle-artefact ('ma') record of the NSTDB will be scaled and added to the test beats at SNR = 18, 6 "
        "and 0 dB, where SNR = 10·log10(P_signal / P_noise), and macro F1 will be plotted against SNR for both "
        "models. Monte Carlo dropout is the planned uncertainty method if time permits; otherwise it is listed "
        "as future work.")

p = find("[TO FILL: programming language")
replace_text(p, "Python 3 with TensorFlow 2.20 / Keras for model building and training; NumPy and pandas for data "
                "handling; scikit-learn for data splits, class weights and metrics; Matplotlib and seaborn for "
                "figures; the wfdb package for reading the PhysioNet MIT-BIH and NSTDB records; the SHAP library "
                "for SHAP attributions; and Streamlit for the demonstration dashboard. Experiments are run on "
                "Kaggle notebooks with an NVIDIA Tesla T4 GPU.")

# ---------- work done till mid-semester --------------------------------------
p_done = find("[TO FILL: completed tasks")
p_prelim = find("[TO FILL: preliminary results")
p_chal = find("[TO FILL: challenges faced]")
for x in (p_prelim, p_chal):
    x._element.getparent().remove(x._element)
replace_text(p_done, "The work completed up to the mid-semester evaluation covers the literature survey, the "
                     "design of the framework and two rounds of preliminary experiments, referred to as V1 "
                     "(baseline) and V2 (step-by-step improvements). Both were run as Kaggle notebooks.")
w = Writer(p_done)
w.sub("Completed Tasks")
w.bullet("Literature survey of deep learning and XAI methods for ECG arrhythmia classification, and "
         "identification of the research gaps (Literature Review chapter).")
w.bullet("Design of the framework: AAMI class mapping, intra- and inter-patient (DS1/DS2) evaluation protocols, "
         "explanation metrics and noise test settings (Methodology chapter).")
w.bullet("V1: a baseline 1D CNN with a class-weighted loss, full test-set evaluation (accuracy, per-class "
         "precision / recall / F1, macro F1, confusion matrix) and Grad-CAM heatmaps for each class.")
w.bullet("V2: five controlled improvement experiments (E1–E5) and validation-based threshold tuning, raising "
         "test macro F1 from 76.29% to 91.10%, followed by Grad-CAM on the final model.")

w.sub("Dataset Used in the Preliminary Experiments")
w.para("To develop and debug the pipeline before building our own pre-processing, the preliminary experiments use "
       "the pre-segmented MIT-BIH heartbeat dataset published by Kachuee et al. [35]. Each beat is 187 samples "
       "long (1.5 s at 125 Hz), scaled to [0, 1] and zero-padded, and is labelled with one of five AAMI classes "
       "(N, S, V, F, Q). The provided training file (87,554 beats) was split 80/20 with stratification "
       "(random seed 42) into training (70,043) and validation (17,511) sets; the provided test file "
       "(21,892 beats) was used only for final evaluation. Table 2 shows the strong class imbalance: a model that "
       "always predicts N would already exceed 80% accuracy, so macro F1 is the main metric. Because this split "
       "is by beat, the same patients appear in training and test data, so all preliminary results are "
       "intra-patient results.")
w.table("Table 2: Class distribution of the dataset used in the preliminary experiments [35]", [
    ["Class", "Training file", "%", "Test file", "%"],
    ["N – Normal", "72,471", "82.77", "18,118", "82.76"],
    ["S – Supraventricular ectopic", "2,223", "2.54", "556", "2.54"],
    ["V – Ventricular ectopic", "5,788", "6.61", "1,448", "6.61"],
    ["F – Fusion", "641", "0.73", "162", "0.74"],
    ["Q – Unknown / paced", "6,431", "7.35", "1,608", "7.35"],
    ["Total", "87,554", "100", "21,892", "100"],
], [2.2, 1.1, 0.7, 1.1, 0.7], bold_rows=(6,))

w.sub("V1: Baseline 1D CNN")
w.para("The baseline is the 1D CNN described in the Methodology chapter (44,741 parameters) with a five-class output. It was "
       "trained with the Adam optimiser (learning rate 10⁻³), batch size 128 and fully balanced class weights "
       "w_c = N / (K·n_c), which give N a weight of 0.24 and F a weight of 27.31. Early stopping monitored "
       "validation loss with a patience of 6 epochs; training stopped after 17 epochs and the weights of "
       "epoch 11 were restored. Table 3 gives the test results and Figure 2 the confusion matrix.")
w.table("Table 3: Test-set results of the V1 baseline 1D CNN", [
    ["Class", "Precision", "Recall", "F1-score", "Test beats"],
    ["N", "0.9813", "0.9594", "0.9702", "18,118"],
    ["S", "0.5586", "0.6691", "0.6088", "556"],
    ["V", "0.8713", "0.8605", "0.8659", "1,448"],
    ["F", "0.2602", "0.8642", "0.4000", "162"],
    ["Q", "0.9896", "0.9502", "0.9695", "1,608"],
    ["Macro average", "0.7322", "0.8607", "0.7629", "21,892"],
], [1.6, 1.1, 1.1, 1.1, 1.1], bold_rows=(6,))
w.figure(FIG / "V1_confusion_matrix.png", "Figure 2: Confusion matrix of the V1 baseline (left: counts; right: "
                                          "row-normalised, diagonal = recall)")
w.para("The baseline reached 94.41% accuracy but only 76.29% macro F1. Class weighting gave high recall on the "
       "rare classes (F recall 86.4%), but at the cost of precision: of the 538 beats predicted as F, only 140 "
       "were truly F (293 were N), giving an F precision of 26.0%. The largest error for S was S → N "
       "(163 of 556 S beats, 29%). Training was also unstable: validation accuracy jumped between 14.5% and "
       "95.1% from one epoch to the next while training loss fell smoothly.")

w.sub("V2: Step-by-step Improvements")
w.para("Each V2 experiment keeps the same data split and seed as V1 and adds one change to the previous "
       "experiment, so that the effect of each change can be seen. The test set was never used to choose an "
       "epoch, a model or a threshold. The changes were: E1, choose the best epoch by validation macro F1 "
       "instead of validation loss; E2, lower learning rate (5×10⁻⁴), gradient clipping (norm 1.0) and batch "
       "normalisation momentum 0.9; E3, square-root class weights (F:N ratio 10.6:1 instead of 113:1); E4, light "
       "training-time augmentation (amplitude scaling ±10%, Gaussian noise σ = 0.01, time shift ±3 samples); and "
       "E5, a residual CNN [36] with three residual blocks (185,477 parameters). Finally, the experiment with the "
       "highest validation macro F1 (E5) was selected and per-class probability scale factors were tuned on the "
       "validation set; the only factor that changed was N × 3. Table 4 and Figure 3 summarise the results.")
w.table("Table 4: Results of the step-by-step experiments (test set, %)", [
    ["ID", "Change", "Acc.", "Macro F1", "S F1", "F F1", "Prec. F", "Rec. F"],
    ["B0", "V1 baseline", "94.41", "76.29", "60.90", "40.00", "26.02", "86.42"],
    ["E1", "Epoch chosen on val. macro F1", "96.28", "83.96", "63.48", "68.12", "55.95", "87.04"],
    ["E2", "+ lower LR, clipping, BN 0.9", "93.58", "78.30", "49.13", "58.49", "43.73", "88.27"],
    ["E3", "+ square-root class weights", "97.39", "87.58", "69.84", "76.97", "75.60", "78.40"],
    ["E4", "+ light augmentation", "97.02", "86.20", "66.16", "76.25", "72.63", "80.25"],
    ["E5", "+ residual network", "98.29", "90.77", "81.22", "78.53", "78.05", "79.01"],
    ["Final", "E5 + tuned thresholds", "98.41", "91.10", "82.14", "79.23", "82.12", "76.54"],
], [0.5, 2.0, 0.6, 0.75, 0.6, 0.6, 0.65, 0.6], bold_rows=(7,))
w.figure(FIG / "V2_08_per_class_f1_progression.png", "Figure 3: Per-class F1 across the V2 experiments",
         width=5.6)
w.para("The largest gains came from three changes: selecting the epoch on validation macro F1 (E1, +7.67 macro "
       "F1), square-root class weights (E3, +9.28 over E2) and the residual network (E5, +4.57 over E4). Two "
       "changes did not help on their own: the training-stability settings of E2 lowered macro F1 by 5.66 points "
       "and did not visibly remove the validation swings, and augmentation (E4) lowered it by 1.38 points. "
       "Threshold tuning added only 0.33 points and traded some S and F recall for precision.")

w.sub("Final Model Results")
w.table("Table 5: Test-set results of the final model (residual CNN + tuned thresholds)", [
    ["Class", "Precision", "Recall", "F1-score", "Test beats"],
    ["N", "0.9874", "0.9962", "0.9918", "18,118"],
    ["S", "0.8758", "0.7734", "0.8214", "556"],
    ["V", "0.9805", "0.9378", "0.9587", "1,448"],
    ["F", "0.8212", "0.7654", "0.7923", "162"],
    ["Q", "0.9981", "0.9832", "0.9906", "1,608"],
    ["Macro average", "0.9326", "0.8912", "0.9110", "21,892"],
], [1.6, 1.1, 1.1, 1.1, 1.1], bold_rows=(6,))
w.figure(FIG / "V2_04_confusion_matrix_final.png", "Figure 4: Confusion matrix of the final model")
w.para("Compared with the baseline, accuracy rose from 94.41% to 98.41% and macro F1 from 76.29% to 91.10%. "
       "F1 improved for every class (S 60.9 → 82.1, V 86.6 → 95.9, F 40.0 → 79.2). False Fusion predictions fell "
       "from 398 to 27 and false S predictions from 294 to 61. Fusion recall, however, is lower than the "
       "baseline's (76.5% vs 86.4%), because the baseline's high recall came from over-predicting F. The largest "
       "remaining error is S → N (120 of 556 S beats, 22%); S beats differ from N beats mainly in timing, which "
       "supports adding explicit RR-interval features in the proposed model.")

w.sub("Grad-CAM Explanations")
w.para("Grad-CAM [37] was applied to the last convolutional activation of each model (46 time steps, "
       "interpolated to the 187-sample beat). For each class, the most confidently and correctly classified test "
       "beat was explained (Figures 5 and 6).")
w.figure(FIG / "V1_gradcam.png", "Figure 5: Grad-CAM heatmaps of the V1 baseline, one beat per class")
w.figure(FIG / "V2_06_gradcam_final.png", "Figure 6: Grad-CAM heatmaps of the final model, one beat per class")
w.para("Both models concentrate their attention on a few narrow regions. The baseline focuses mostly on the start "
       "of the window and on the second large peak or the end of the signal; the final model's maps are broader "
       "and more often cover waveform features such as bumps and troughs. Since each window starts at the "
       "current R peak, the second peak is likely the next R peak, which suggests that the models partly use "
       "the timing of the next beat (an implicit RR cue) as well as beat shape. These observations are "
       "qualitative only; they will be tested with the region-share and deletion measures defined in the Methodology chapter.")

w.sub("Challenges Faced")
w.bullet("Severe class imbalance: fully balanced class weights made the baseline over-predict the rare classes "
         "(F precision 26.0%). Softer, square-root weights fixed most of this.")
w.bullet("Unstable training: the baseline's validation accuracy swung between 14.5% and 95.1%, so early stopping "
         "on validation loss stopped training too early. Choosing the epoch on validation macro F1 and the "
         "residual network gave smoother, better results.")
w.bullet("Not every planned improvement helped: the stability settings (E2) and augmentation (E4) lowered macro "
         "F1, and because the experiments were cumulative, the best combination of changes is not yet known.")
w.bullet("The preliminary data are pre-processed, intra-patient and include the Q class, so the results are "
         "optimistic and not directly comparable with the planned four-class, inter-patient results. Each "
         "experiment was also run only once with one seed.")
w.bullet("Grad-CAM is coarse (46 time steps) and has so far been shown only for one hand-picked beat per class.")

# ---------- work plan ---------------------------------------------------------
p = find("The following tasks follow from the project objectives.")
replace_text(p, "The following tasks follow from the project objectives and build on the pipeline developed so "
                "far. Table 6 gives the planned timeline.")
w = Writer(p)
w.bullet("Port the pipeline from the pre-segmented Kaggle beats to the raw MIT-BIH records (own filtering, "
         "252-sample beats around the R peak, RR features, four classes N, S, V, F).")
last = find("Analyse results, prepare the final report and presentation.")
w = Writer(last)
w.table("Table 6: Planned timeline for the next phase", [
    ["Period", "Task", "Output"],
    ["12 – 25 Oct", "Raw MIT-BIH pre-processing; re-train the baseline CNN on 4 classes", "Beat dataset, class counts, baseline results"],
    ["26 Oct – 1 Nov", "CNN-LSTM + RR model; DS1/DS2 inter-patient run; SHAP", "Model comparison table, Grad-CAM vs SHAP figure"],
    ["2 – 8 Nov", "Region-share and deletion tests", "Region table, deletion chart"],
    ["9 – 15 Nov", "NSTDB noise at 18 / 6 / 0 dB; heatmap stability", "F1-vs-SNR and stability plots"],
    ["16 – 22 Nov", "Streamlit demonstration dashboard", "Working demo"],
    ["23 – 29 Nov", "Buffer; tidy code; freeze results", "Final figures and tables"],
    ["30 Nov – 27 Dec", "Final report writing and supervisor review", "Final report"],
    ["28 Dec – 10 Jan", "Presentation and viva preparation", "Final slides"],
], [1.3, 2.7, 2.0])

# ---------- expected results ------------------------------------------------
p = find("[TO FILL: target performance values")
replace_text(p, "On the intra-patient split, macro F1 of about 90% is expected, in line with the 91.10% already "
                "reached in the preliminary experiments. On the inter-patient DS1/DS2 split, performance is "
                "expected to be clearly lower, as reported in the literature [1], [11], with S and F remaining "
                "the hardest classes; the RR features are expected to improve S recall in particular. For the "
                "explanations, V beats are expected to place a larger share of importance on the QRS region than "
                "N beats, and deleting the most important samples should lower model confidence much more than "
                "deleting random samples. Under NSTDB noise, macro F1 and explanation stability are expected to "
                "fall as the SNR decreases from 18 dB to 0 dB.")

# ---------- references --------------------------------------------------------
last_ref = find("[34] K. Mallikarjunamallu")
w = Writer(last_ref)
for ref in [
    "[35] M. Kachuee, S. Fazeli, and M. Sarrafzadeh, \"ECG heartbeat classification: A deep transferable "
    "representation,\" in Proc. IEEE Int. Conf. Healthcare Informatics (ICHI), 2018, pp. 443–444.",
    "[36] K. He, X. Zhang, S. Ren, and J. Sun, \"Deep residual learning for image recognition,\" in Proc. IEEE "
    "Conf. Computer Vision and Pattern Recognition (CVPR), 2016, pp. 770–778.",
    "[37] R. R. Selvaraju, M. Cogswell, A. Das, R. Vedantam, D. Parikh, and D. Batra, \"Grad-CAM: Visual "
    "explanations from deep networks via gradient-based localization,\" in Proc. IEEE Int. Conf. Computer "
    "Vision (ICCV), 2017, pp. 618–626.",
    "[38] S. M. Lundberg and S.-I. Lee, \"A unified approach to interpreting model predictions,\" in Proc. "
    "Advances in Neural Information Processing Systems (NeurIPS), 2017, pp. 4765–4774.",
]:
    new = copy.deepcopy(last_ref._element)
    w.cur.addnext(new)
    w.cur = new
    pp = Paragraph(new, d._body)
    for r in pp.runs[1:]:
        r._element.getparent().remove(r._element)
    pp.runs[0].text = ref

d.save(str(DOC))
left = [p.text for p in d.paragraphs if "TO FILL" in p.text]
print("saved; remaining TO FILL:", left)
