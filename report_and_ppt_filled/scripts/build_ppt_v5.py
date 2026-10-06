"""Build the final 15-slide mid-sem deck (v5) by condensing MidSem_PPT_ECG_XAI_v4.pptx.

Usage: python build_ppt_v5.py <v4.pptx> <out_v5.pptx>
V4 is only read. Its title slide, bars, title boxes, fonts, card/table styles and figures are reused
with the same helpers that built v4; all content comes from the v4 slides. References are cut to the
key ones and renumbered [1]..[n] in order of first citation.
"""
import copy
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SRC, OUT = Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve()
assert SRC != OUT
FIG = SRC.parent / "figures"
prs = Presentation(str(SRC))
DARK, GREY, NAVY, ORANGE, GREEN, BLUE = "202020", "595959", "1F3864", "ED7D31", "548235", "2E75B6"
TABLE_STYLE = "{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}"


# ---------- v4 style helpers (identical to build_ppt_v4.py) --------------------
def is_bar(sh):
    return sh.width == prs.slide_width and (sh.top < Emu(200000) or sh.top > Emu(6600000))


def is_title(sh):
    return (not is_bar(sh) and sh.has_text_frame and sh.text_frame.text.strip() != ""
            and sh.top < Emu(500000) and sh.width > Emu(4000000))


def set_text(shape, text):
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    p0.runs[0].text = text


def textbox(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, space_after=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    bp = tf._txBody.find(qn("a:bodyPr"))
    for k in ("lIns", "tIns", "rIns", "bIns"):
        bp.set(k, "0")
    if bp.find(qn("a:spAutoFit")) is None:  # a second spAutoFit makes PowerPoint reject the file
        etree.SubElement(bp, qn("a:spAutoFit"))
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space_after:
            p.space_after = Pt(space_after)
        for text, size, bold, color, italic in (para if isinstance(para, list) else [para]):
            r = p.add_run()
            r.text = text
            f = r.font
            f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, "Calibri"
            f.color.rgb = RGBColor.from_string(color)
    return tb


def heading(slide, x, y, w, text, size=15, color=DARK):
    return textbox(slide, x, y, w, 0.35, [(text, size, True, color, False)])


def body(slide, x, y, w, h, lines, size=12, color=DARK, space_after=3):
    paras = []
    for ln in lines:
        if isinstance(ln, list):
            paras.append([(t, size, b, color, False) for t, b in ln])
        else:
            paras.append((ln, size, False, color, False))
    return textbox(slide, x, y, w, h, paras, space_after=space_after)


def note(slide, x, y, w, text, size=11, align=PP_ALIGN.LEFT):
    return textbox(slide, x, y, w, 0.3, [(text, size, False, GREY, True)], align=align)


def accent(slide, x, y, h, color, w=0.07):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(color)
    sh.line.fill.background()
    return sh


def card(slide, x, y, w, h, title, text, color, size=11.5):
    accent(slide, x, y, h, color)
    heading(slide, x + 0.18, y + 0.05, w - 0.2, title, size=13, color=color)
    lines = text if isinstance(text, list) else [text]
    body(slide, x + 0.18, y + 0.42, w - 0.2, h - 0.45, lines, size=size)


def table(slide, x, y, w, rows, col_w, row_h=0.27, size=11, hl_rows=(), left_cols=(0,)):
    gf = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                Inches(row_h * len(rows)))
    tbl = gf.table
    tbl._tbl.tblPr.find(qn("a:tableStyleId")).text = TABLE_STYLE
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.margin_left = c.margin_right = Inches(0.06)
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = RGBColor.from_string(NAVY)
            elif i in hl_rows:
                c.fill.fore_color.rgb = RGBColor.from_string("FCE4D6")
            else:
                c.fill.fore_color.rgb = RGBColor.from_string("F2F2F2" if i % 2 else "FFFFFF")
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j in left_cols else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(val)
            f = r.font
            f.size, f.name = Pt(size), "Calibri"
            f.bold = i == 0 or j == 0 or i in hl_rows
            f.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else DARK)
    return gf


def picture(slide, name, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(str(FIG / name), Inches(x), Inches(y), **kw)


def caption(slide, x, y, w, text):
    return note(slide, x, y, w, text, size=10, align=PP_ALIGN.CENTER)


def numbered(slide, x, y, w, items, step, size=12.5, num_size=18):
    """v4 'Work Plan' style: orange number + navy bar + text."""
    for k, t in enumerate(items):
        yy = y + k * step
        textbox(slide, x, yy, 0.4, 0.4, [(f"{k + 1}", num_size, True, ORANGE, False)])
        accent(slide, x + 0.4, yy + 0.03, step - 0.15, NAVY, w=0.05)
        if isinstance(t, list):
            body(slide, x + 0.55, yy + 0.04, w - 0.55, step, [t], size=size)
        else:
            body(slide, x + 0.55, yy + 0.04, w - 0.55, step, [t], size=size)


TEMPLATE = prs.slides[1]  # a v4 content slide: bars + title box
BLANK = TEMPLATE.slide_layout
OLD = list(prs.slides._sldIdLst)[1:]  # v4 slides 2-28, removed at the end


def new_slide(title):
    s = prs.slides.add_slide(BLANK)
    for sh in TEMPLATE.shapes:
        if is_bar(sh) or is_title(sh):
            s.shapes._spTree.append(copy.deepcopy(sh._element))
    for sh in s.shapes:
        if is_title(sh):
            set_text(sh, title)
    return s


# ---------- references: v4/report numbers -> new [1..n] in order of first use ----
REFS = {
    1: "Q. Xiao et al., \"Deep learning-based ECG arrhythmia classification: A systematic review,\" "
       "Applied Sciences, vol. 13, Art. no. 4964, 2023.",
    2: "N. N. Kulkarni et al., \"Recent advances in deep learning based arrhythmia classification using ECG "
       "and PPG signals: A systematic review,\" Discover Computing, vol. 29, Art. no. 666, 2026.",
    3: "J. Beck and A. John, \"Explainable AI (XAI) for arrhythmia detection from electrocardiograms,\" "
       "arXiv:2508.17294, 2025.",
    5: "B. A. Hassoon et al., \"Dual-branch bidirectional attention fusion of temporal and hierarchical "
       "representations for robust ECG arrhythmia classification,\" Expert Systems with Applications, "
       "vol. 331, Art. no. 133409, 2026.",
    8: "N. Alamatsaz et al., \"A lightweight hybrid CNN-LSTM explainable model for ECG-based arrhythmia "
       "detection,\" Biomedical Signal Processing and Control, vol. 90, Art. no. 105884, 2024.",
    11: "Sa. I. Ibrahim, \"A hybrid deep learning algorithm for ECG-based heart disease classification,\" "
        "Scientific Reports, vol. 16, Art. no. 29911, 2026.",
    35: "M. Kachuee, S. Fazeli, and M. Sarrafzadeh, \"ECG heartbeat classification: A deep transferable "
        "representation,\" Proc. IEEE ICHI, 2018, pp. 443–444.",
    37: "R. R. Selvaraju et al., \"Grad-CAM: Visual explanations from deep networks via gradient-based "
        "localization,\" Proc. IEEE ICCV, 2017, pp. 618–626.",
    38: "S. M. Lundberg and S.-I. Lee, \"A unified approach to interpreting model predictions,\" Proc. "
        "NeurIPS, 2017, pp. 4765–4774.",
}
ORDER = []


def c(*keys):
    """Citation text for v4 reference numbers, renumbered by first use."""
    out = []
    for k in keys:
        assert k in REFS, k
        if k not in ORDER:
            ORDER.append(k)
        out.append(f"[{ORDER.index(k) + 1}]")
    return ", ".join(out)


# =================== 2. Introduction & Motivation ===================
s = new_slide("Introduction & Motivation")
heading(s, 0.55, 1.2, 8.9, "Background")
body(s, 0.55, 1.62, 8.9, 1.6, [
    "• The ECG is a widely used non-invasive tool for detecting cardiac arrhythmias; automating its "
    "interpretation is an active research area.",
    f"• Deep learning dominates: CNNs are used in 58.7% of 368 studies {c(1)} and 62% of 119 studies "
    f"{c(2)}; hybrid CNN-recurrent and Transformer models are increasingly common.",
], size=13)
heading(s, 0.55, 3.25, 8.9, "Motivation")
card(s, 0.55, 3.75, 4.35, 1.5, "Accuracy that does not transfer",
     f"Reported accuracies above 98% fall considerably on unseen patients, external databases or noisy "
     f"signals {c(1, 2)}.", NAVY, size=12.5)
card(s, 5.1, 3.75, 4.35, 1.5, "Explanations that are not checked",
     f"Explanations attached to such models are rarely validated against clinical ECG knowledge {c(3)}.",
     ORANGE, size=12.5)
note(s, 0.55, 5.6, 8.9, "Clinical use needs a classifier that works on new patients, explains its decisions, "
                          "and stays reliable on real, noisy recordings.", size=12)

# =================== 3. Problem Statement ===================
s = new_slide("Problem Statement")
accent(s, 0.55, 1.2, 1.05, ORANGE)
body(s, 0.75, 1.25, 8.7, 1.0, [
    "There is a need for an ECG arrhythmia classification framework whose performance is measured on "
    "unseen patients, whose explanations are checked for consistency and clinical plausibility, and whose "
    "predictions and explanations are evaluated for robustness under realistic noise."], size=14)
heading(s, 0.55, 2.6, 8.9, "Existing limitations")
card(s, 0.55, 3.1, 2.85, 1.65, "Evaluation protocol",
     f"F1 fell from 95.52% (intra-patient) to 83.89% (inter-patient) in studies reporting both {c(1)}.",
     NAVY, size=12)
card(s, 3.6, 3.1, 2.85, 1.65, "Black-box models",
     f"Fewer than 10% of studies validated explanation maps against clinically meaningful ECG structures "
     f"{c(2)}.", ORANGE, size=12)
card(s, 6.65, 3.1, 2.85, 1.65, "Narrow robustness tests",
     f"Noise robustness is often tested with additive white Gaussian noise only {c(5)}.", GREEN, size=12)
body(s, 0.55, 5.1, 8.9, 1.0, [
    [("Why it matters: ", True), ("a model that is only tested on familiar patients and clean signals, with "
                                 "unchecked explanations, cannot be trusted on new, real recordings.", False)],
], size=12.5)

# =================== 4. Objectives ===================
s = new_slide("Objectives")
objs = [
    ("Patient-independent baseline",
     "Develop a hybrid deep learning arrhythmia classifier evaluated under a strict inter-patient protocol "
     "(MIT-BIH DS1/DS2) with external testing on a second database; report per-class and macro F1 "
     "alongside accuracy."),
    ("Validated explanations",
     "Apply post-hoc XAI methods (SHAP and Grad-CAM) and measure the within-class consistency of "
     "explanations and their agreement with ECG features (P wave, QRS complex, RR interval)."),
    ("Robustness of predictions and explanations",
     "Assess how classification performance and explanation stability degrade under realistic NSTDB noise "
     "at graded SNRs, using uncertainty estimates to flag unreliable predictions."),
]
for k, (t, d) in enumerate(objs):
    y = 1.3 + k * 1.8
    textbox(s, 0.55, y, 0.7, 0.5, [(f"0{k + 1}", 26, True, ORANGE, False)])
    heading(s, 1.35, y, 8.1, f"Objective {k + 1}: {t}", size=15, color=NAVY)
    body(s, 1.35, y + 0.45, 8.1, 1.1, [d], size=13)

# =================== 5. Literature Review ===================
s = new_slide("Literature Review")
note(s, 0.55, 1.12, 8.9, "36 papers surveyed (5 reviews, 31 original studies). Most relevant studies:", size=12)
table(s, 0.55, 1.55, 8.9, [
    ["Ref.", "Dataset(s)", "Model", "XAI", "Split", "Key result"],
    [c(8), "MIT-BIH + LTAF", "Lightweight CNN-LSTM", "SHAP", "Not stated (85/15)",
     "Acc 98.24%, Se 86.1%; 0.16 MB"],
    [c(3), "MIT-BIH + 12-lead", "1D CNN", "4 SHAP variants", "Not stated (80/20)",
     "Val acc 98.30%; 73.45% on combined data"],
    [c(5), "MIT-BIH; INCART, SVDB, PTB", "CNN-BiLSTM + CapsNet fusion", "Branch weights", "Inter-patient",
     "Acc 99.55%; F1 32.46% at 5 dB AWGN"],
    [c(11), "MIT-BIH", "CNN-BiLSTM + cGAN", "Not stated", "Both", "99.00% beat-wise vs 91.69% patient-wise"],
], [0.55, 1.75, 1.75, 1.15, 1.3, 2.4], row_h=0.5, size=11, left_cols=(0, 1, 2, 5))
heading(s, 0.55, 4.25, 8.9, "Key findings", size=14)
body(s, 0.55, 4.65, 8.9, 2.3, [
    f"• Architectures: CNNs dominate; hybrids add a recurrent or attention module (CNN-LSTM {c(8)}, "
    f"CNN-BiLSTM {c(11)}); Transformers gain 0.845% over 1D-CNNs at higher cost {c(2)}.",
    f"• XAI: SHAP is the most used method, then Grad-CAM; APB/PVC explanations were inconsistent within a "
    f"class {c(3)}.",
    f"• Evaluation: MIT-BIH dominates (61% of studies {c(1)}); external validation in only 18.5% of studies, "
    f"with a mean accuracy drop of 8.7% {c(2)}.",
], size=12)

# =================== 6. Research Gap ===================
s = new_slide("Research Gap")
gaps = [
    ("Gap 1 – Explanations rarely validated", NAVY,
     f"< 10% of studies validate maps against ECG structures {c(2)}; explanations inconsistent within a "
     f"class {c(3)}.", "Objective 2: region-share and deletion checks"),
    ("Gap 2 – XAI seldom patient-independent", ORANGE,
     f"XAI studies often do not describe patient separation, though inter-patient evaluation lowers "
     f"performance {c(1, 11)}.", "Objectives 1 + 2: XAI on a DS1/DS2 model"),
    ("Gap 3 – Noise tests narrow, not linked to XAI", GREEN,
     f"Noise tests limited to AWGN {c(5)} or absent; no surveyed paper tests whether explanations stay "
     f"stable under noise.", "Objective 3: NSTDB noise + heatmap stability"),
    ("Gap 4 – Weak minority classes", BLUE,
     f"Patient-wise Fusion and S recall only 0.021 and 0.209 {c(11)}; little external validation {c(2)}.",
     "Objective 1: per-class / macro F1, external test"),
]
for k, (t, col, d, fix) in enumerate(gaps):
    x = 0.55 if k % 2 == 0 else 5.1
    y = 1.4 if k < 2 else 3.95
    card(s, x, y, 4.35, 2.1, t, [d, [("→ Addressed by ", True), (fix, False)]], col, size=12)

# =================== 7. Dataset ===================
s = new_slide("Dataset")
heading(s, 0.55, 1.12, 4.35, "Framework datasets", size=14)
body(s, 0.55, 1.5, 4.35, 1.6, [
    "• MIT-BIH Arrhythmia Database: AAMI classes, inter-patient DS1/DS2 split.",
    "• External test: INCART (preferred; stretch goal).",
    "• MIT-BIH NSTDB: baseline wander, muscle artefact, electrode motion noise.",
], size=11.5)
heading(s, 5.1, 1.12, 4.35, "Used so far (preliminary)", size=14)
body(s, 5.1, 1.5, 4.35, 1.6, [
    f"• Pre-segmented MIT-BIH heartbeat dataset {c(35)}.",
    "• 187-sample beats (1.5 s at 125 Hz), scaled to [0, 1], zero-padded.",
    "• 5 AAMI classes: N, S, V, F, Q.",
], size=11.5)
table(s, 0.55, 3.1, 8.9, [
    ["File", "N", "S", "V", "F", "Q", "Total"],
    ["Train file", "72,471 (82.77%)", "2,223 (2.54%)", "5,788 (6.61%)", "641 (0.73%)", "6,431 (7.35%)",
     "87,554"],
    ["Test file", "18,118 (82.76%)", "556 (2.54%)", "1,448 (6.61%)", "162 (0.74%)", "1,608 (7.35%)", "21,892"],
], [1.1, 1.45, 1.2, 1.25, 1.15, 1.25, 1.0], row_h=0.3, size=10.5)
note(s, 0.55, 4.05, 8.9, "Split: train file 80/20 (stratified, seed 42) → 70,043 train / 17,511 validation; "
                          "test file (21,892) used only for final evaluation. Split is by beat → intra-patient.",
     size=10.5)
picture(s, "V1_example_beats.png", 0.55, 4.65, w=8.9)
caption(s, 0.55, 6.15, 8.9, "One example beat per class: starts at the R peak, second peak ≈ next R peak, then "
                            "zero padding")

# =================== 8. Proposed Methodology ===================
s = new_slide("Proposed Methodology")
picture(s, "block_diagram.png", 1.75, 1.05, w=6.5)
caption(s, 0.55, 5.15, 8.9, "Block diagram of the proposed framework")
body(s, 0.55, 5.5, 4.35, 1.4, [
    [("Pre-processing: ", True), ("filtering, R-peak-centred beats, normalisation, RR features.", False)],
    [("Split: ", True), ("intra-patient 80/20 + inter-patient DS1/DS2.", False)],
], size=11.5)
body(s, 5.1, 5.5, 4.35, 1.4, [
    [("Explain: ", True), (f"Grad-CAM {c(37)} + Gradient SHAP {c(38)}.", False)],
    [("Robustness: ", True), ("NSTDB 'ma' noise at 18 / 6 / 0 dB SNR; macro F1 vs SNR.", False)],
], size=11.5)

# =================== 9. Model / Algorithms ===================
s = new_slide("Model & Algorithms")
heading(s, 0.55, 1.1, 8.9, "Baseline 1D CNN (implemented): input 187 × 1 beat → softmax over 5 classes", size=13)
table(s, 0.55, 1.5, 8.9, [
    ["Layer", "Output shape", "Parameters"],
    ["Conv1D 32, k = 7 → BN → ReLU → MaxPool 2", "(93, 32)", "256 + 128"],
    ["Conv1D 64, k = 5 → BN → ReLU → MaxPool 2", "(46, 64)", "10,304 + 256"],
    ["Conv1D 128, k = 3 → BN → ReLU (Grad-CAM layer) → MaxPool 2", "(23, 128)", "24,704 + 512"],
    ["Global average pooling → Dense 64, ReLU → Dropout 0.3 → Dense 5, softmax", "(5)", "8,256 + 325"],
    ["Total", "", "44,741"],
], [5.6, 1.4, 1.9], row_h=0.27, size=11, hl_rows=(5,))
card(s, 0.55, 3.4, 2.85, 2.0, "Residual CNN (V2 final)",
     "Conv 32 (k 7) stem + 3 residual blocks (32, 64, 128); each: two Conv1D (k 5) + BN with a shortcut "
     "(1×1 conv when channels change). 185,477 parameters.", NAVY, size=11)
card(s, 3.6, 3.4, 2.85, 2.0, "Proposed: CNN-LSTM + RR",
     "Same conv front end → LSTM (64); last hidden state + previous / next RR (÷ record mean RR) → "
     "Dense 64 → softmax over N, S, V, F.", ORANGE, size=11)
card(s, 6.65, 3.4, 2.85, 2.0, "Explainability",
     f"Grad-CAM {c(37)} on the last conv activation (46 steps → 187 samples). Gradient SHAP {c(38)} planned. "
     "Measures: region share, deletion test, stability.", GREEN, size=11)
body(s, 0.55, 5.75, 8.9, 0.6, [
    [("Key design choices: ", True), ("square-root class weights (F:N 10.6:1 instead of 113:1); best epoch "
                                     "chosen on validation macro F1; per-class probability thresholds tuned "
                                     "on validation.", False)]], size=11.5)

# =================== 10. Experimental Setup ===================
s = new_slide("Experimental Setup")
table(s, 0.55, 1.1, 8.9, [
    ["ID", "Change added (cumulative)", "Network", "LR", "Clip norm", "BN mom.", "Class weights", "Augment."],
    ["E1", "Epoch chosen on val. macro F1", "CNN", "1×10⁻³", "–", "0.99", "Balanced", "No"],
    ["E2", "+ lower LR, clipping, BN 0.9", "CNN", "5×10⁻⁴", "1.0", "0.9", "Balanced", "No"],
    ["E3", "+ square-root class weights", "CNN", "5×10⁻⁴", "1.0", "0.9", "Square-root", "No"],
    ["E4", "+ light augmentation", "CNN", "5×10⁻⁴", "1.0", "0.9", "Square-root", "Yes"],
    ["E5", "+ residual network", "Residual CNN", "5×10⁻⁴", "1.0", "0.9", "Square-root", "Yes"],
], [0.5, 2.45, 1.15, 0.8, 0.8, 0.8, 1.2, 0.85], row_h=0.27, size=10.5, left_cols=(0, 1))
note(s, 0.55, 2.8, 8.9, "B0 = V1 baseline (LR 10⁻³, balanced weights, early stopping on validation loss, "
                         "patience 6). Final = E5 + thresholds tuned on validation (N × 3).", size=10.5)
card(s, 0.55, 3.4, 2.85, 2.6, "Training", [
    "Adam, batch 128, ≤ 30 epochs",
    "Best weights on val. macro F1; LR halved after 3 non-improving epochs (min 10⁻⁵); stop after 8",
    "Augmentation: amplitude ±10%, noise σ 0.01, shift ±3 samples",
    "Seed 42; test set never used for selection",
], BLUE, size=11)
card(s, 3.6, 3.4, 2.85, 2.6, "Evaluation metrics", [
    "Accuracy",
    "Per-class precision, recall, F1",
    "Macro F1 – main metric (N ≈ 83% of beats, so accuracy hides rare-class errors)",
    "Confusion matrix",
], ORANGE, size=11)
card(s, 6.65, 3.4, 2.85, 2.6, "Tools & hardware", [
    "Python 3, TensorFlow 2.20 / Keras",
    "NumPy, pandas, scikit-learn, Matplotlib, seaborn",
    "wfdb, SHAP, Streamlit (next phase)",
    "Kaggle notebooks, NVIDIA Tesla T4 GPU; 5 runs in 8.1 min",
], GREEN, size=11)

# =================== 11. Work Done Till Now ===================
s = new_slide("Work Done Till Now")
heading(s, 0.55, 1.1, 4.4, "Completed milestones", size=14)
numbered(s, 0.55, 1.55, 4.4, [
    "Literature survey of 36 papers; research gaps identified",
    "Framework design: AAMI classes, DS1/DS2 protocol, XAI metrics, noise settings",
    "Data preparation: stratified 80/20 split, class weights",
    "V1: baseline 1D CNN trained, evaluated, Grad-CAM per class",
    "V2: 5 controlled experiments (E1–E5) + threshold tuning",
    "Final model evaluated; Grad-CAM on the final model",
], step=0.83, size=11.5, num_size=16)
heading(s, 5.15, 1.1, 4.3, "Experiments conducted (test set, %)", size=14)
table(s, 5.15, 1.55, 4.3, [
    ["ID", "Change", "Acc.", "Macro F1"],
    ["B0", "V1 baseline", "94.41", "76.29"],
    ["E1", "Epoch on val. macro F1", "96.28", "83.96"],
    ["E2", "+ LR, clipping, BN", "93.58", "78.30"],
    ["E3", "+ sqrt class weights", "97.39", "87.58"],
    ["E4", "+ augmentation", "97.02", "86.20"],
    ["E5", "+ residual network", "98.29", "90.77"],
    ["Final", "E5 + tuned thresholds", "98.41", "91.10"],
], [0.6, 2.0, 0.75, 0.95], row_h=0.3, size=10.5, hl_rows=(7,), left_cols=(0, 1))
note(s, 5.15, 4.05, 4.3, "Same data, split and seed for every step; one change added at a time.", size=10.5)
picture(s, "V2_02_validation_curves.png", 5.15, 4.5, w=4.3)
caption(s, 5.15, 5.75, 4.3, "Validation macro F1 / accuracy per epoch, E1–E5")

# =================== 12. Preliminary Results ===================
s = new_slide("Preliminary Results")
textbox(s, 0.55, 1.08, 8.9, 0.4, [[
    ("Accuracy 94.41% → ", 15, True, NAVY, False), ("98.41%", 15, True, ORANGE, False),
    ("     |     Macro F1 76.29% → ", 15, True, NAVY, False), ("91.10%", 15, True, ORANGE, False),
    ("   (V1 baseline → V2 final)", 12, False, GREY, True)]])
picture(s, "V2_04_confusion_matrix_final.png", 0.55, 1.55, w=5.3)
caption(s, 0.55, 3.62, 5.3, "Final model confusion matrix (counts; row-normalised)")
picture(s, "V2_08_per_class_f1_progression.png", 6.0, 1.55, w=3.45)
caption(s, 6.0, 3.42, 3.45, "Per-class F1, B0 → Final")
table(s, 0.55, 3.95, 5.3, [
    ["Metric", "V1 baseline", "V2 final"],
    ["Precision / Recall S (%)", "55.86 / 66.91", "87.58 / 77.34"],
    ["Precision / Recall F (%)", "26.02 / 86.42", "82.12 / 76.54"],
    ["Beats wrongly predicted as F", "398", "27"],
    ["Beats wrongly predicted as S", "294", "61"],
    ["S beats predicted as N", "163 (29%)", "120 (22%)"],
    ["Parameters", "44,741", "185,477"],
], [2.5, 1.4, 1.4], row_h=0.28, size=10.5)
table(s, 6.0, 3.95, 3.45, [
    ["Class", "Prec.", "Rec.", "F1"],
    ["N", "0.987", "0.996", "0.992"],
    ["S", "0.876", "0.773", "0.821"],
    ["V", "0.981", "0.938", "0.959"],
    ["F", "0.821", "0.765", "0.792"],
    ["Q", "0.998", "0.983", "0.991"],
    ["Macro", "0.933", "0.891", "0.911"],
], [0.75, 0.9, 0.9, 0.9], row_h=0.28, size=10.5, hl_rows=(6,))
note(s, 0.55, 6.0, 8.9, "Test set: 21,892 beats; intra-patient split. F1 improved for every class (S 60.9 → 82.1, "
                         "V 86.6 → 95.9, F 40.0 → 79.2).", size=10.5)

# =================== 13. Challenges & Key Findings ===================
s = new_slide("Challenges & Key Findings")
heading(s, 0.55, 1.1, 4.35, "Challenges faced", size=14, color=ORANGE)
body(s, 0.55, 1.5, 4.35, 2.6, [
    "• Class imbalance: balanced weights (F:N 113:1) gave F precision 26.0%.",
    "• Unstable training: V1 validation accuracy swung 14.5% ↔ 95.1%.",
    "• E2 stability settings (−5.66) and augmentation E4 (−1.38) hurt; steps cumulative.",
    "• Optimistic setting: pre-processed, intra-patient, 5 classes, one seed per run.",
], size=11.5)
heading(s, 5.1, 1.1, 4.35, "Key findings", size=14, color=GREEN)
body(s, 5.1, 1.5, 4.35, 2.6, [
    "• Biggest gains: square-root weights (E3, +9.28), epoch selection on macro F1 (E1, +7.67), "
    "residual network (E5, +4.57).",
    "• F recall fell 86.4% → 76.5%: V1's high recall came from over-predicting F.",
    "• Largest remaining error S → N (22%): S differs from N mainly in timing → RR features.",
], size=11.5)
heading(s, 0.55, 3.55, 8.9, "Explainability finding", size=14, color=NAVY)
picture(s, "V2_06_gradcam_final.png", 0.55, 3.95, w=8.9)
caption(s, 0.55, 5.85, 8.9, "Grad-CAM of the final model (most confident correct beat per class): heat often near "
                            "the second R peak → possible implicit RR cue; qualitative only so far")

# =================== 14. Future Work & Expected Outcomes ===================
s = new_slide("Future Work & Expected Outcomes")
heading(s, 0.55, 1.1, 4.35, "Remaining tasks", size=14)
numbered(s, 0.55, 1.5, 4.35, [
    "Port pipeline to raw MIT-BIH; pre-processing + DS1/DS2 split",
    "Train hybrid CNN-LSTM + RR; per-class and macro F1; external test",
    "Apply Grad-CAM + SHAP; region-share, deletion and consistency checks",
    "NSTDB noise at graded SNRs; explanation stability and uncertainty",
    "Analyse results; final report and presentation",
], step=0.66, size=11, num_size=14)
heading(s, 0.55, 4.95, 4.35, "Anticipated challenges", size=14, color=ORANGE)
body(s, 0.55, 5.33, 4.35, 1.5, [
    "• Inter-patient scores clearly lower; S and F remain hardest.",
    "• External testing and MC dropout only if time permits.",
], size=11)
heading(s, 5.1, 1.1, 4.35, "Deliverables", size=14)
body(s, 5.1, 1.5, 4.35, 2.0, [
    "• Patient-independent classifier with per-class and macro results (inter-patient + external).",
    "• Quantitative evaluation of explanations with two XAI methods.",
    "• Robustness and uncertainty analysis under NSTDB noise.",
    "• Final report and presentation.",
], size=11)
heading(s, 5.1, 3.35, 4.35, "Expected results", size=14, color=GREEN)
body(s, 5.1, 3.75, 4.35, 2.5, [
    f"• Intra-patient macro F1 ≈ 90%; DS1/DS2 clearly lower {c(1, 11)}; RR features should raise S recall.",
    "• V beats: larger QRS importance share than N; top-sample deletion ≫ random deletion.",
    "• Macro F1 and explanation stability fall as SNR drops 18 → 0 dB.",
], size=11)

# =================== 15. Conclusion & References ===================
s = new_slide("Conclusion & References")
heading(s, 0.55, 1.1, 8.9, "Conclusion", size=14)
body(s, 0.55, 1.5, 8.9, 2.4, [
    "• A working training, evaluation and Grad-CAM pipeline is in place.",
    [("• Macro F1 ", False), ("76.29% → 91.10%", True), (" and accuracy ", False), ("94.41% → 98.41%", True),
     (" via epoch selection, square-root weights and a residual CNN.", False)],
    "• Results are intra-patient: a development baseline for the patient-independent framework.",
    [("• Next: ", True), ("raw MIT-BIH + DS1/DS2 → CNN-LSTM + RR → Grad-CAM / SHAP checks → NSTDB noise.", False)],
], size=12.5)
heading(s, 0.55, 3.45, 8.9, "References", size=14)
ref_box = body(s, 0.55, 3.85, 8.9, 3.0, ["placeholder"], size=10.5, space_after=3)

# ---------- fill references in order of first citation ------------------------
tf = ref_box.text_frame
tmpl = copy.deepcopy(tf.paragraphs[0]._p)
for p in list(tf.paragraphs):
    p._p.getparent().remove(p._p)
for i, k in enumerate(ORDER, 1):
    np_ = copy.deepcopy(tmpl)
    np_.find(qn("a:r")).find(qn("a:t")).text = f"[{i}] {REFS[k]}"
    tf._txBody.append(np_)

# ---------- remove v4 slides 2-28 --------------------------------------------
ids = prs.slides._sldIdLst
for el in OLD:
    prs.part.drop_rel(el.get(qn("r:id")))
    ids.remove(el)

# ---------- checks -------------------------------------------------------------
assert len(prs.slides) <= 15, len(prs.slides)
used = set()
for sl in prs.slides:
    for sh in sl.shapes:
        texts = [sh.text_frame.text] if sh.has_text_frame else []
        if getattr(sh, "has_table", False) and sh.has_table:
            texts += [cl.text for r in sh.table.rows for cl in r.cells]
        for t in texts:
            if not t.startswith("[1] "):
                used |= {int(n) for n in re.findall(r"\[(\d+)\]", t)}
assert used == set(range(1, len(ORDER) + 1)), (used, len(ORDER))
prs.save(str(OUT))
print("saved", OUT.name, "| slides:", len(prs.slides), "| references:", len(ORDER))
