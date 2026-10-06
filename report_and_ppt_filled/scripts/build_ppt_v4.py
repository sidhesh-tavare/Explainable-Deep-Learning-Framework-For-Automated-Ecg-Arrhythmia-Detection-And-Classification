"""Build MidSem_PPT_ECG_XAI_v4.pptx from the v3 deck, with every content slide rewritten strictly
from UG_Project_Report_ECG_XAI_filled_v2.docx (text, numbers, figures, tables and reference numbers).

Usage: python build_ppt_v4.py <v3.pptx> <out_v4.pptx>
The v3 deck is only read. Slide 1 (title) and the v3 visual style (bars, title box, fonts, table
colours) are kept; slides 2-18 of v3 are replaced.
"""
import copy
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SRC, OUT = Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve()
assert SRC != OUT
FIG = SRC.parent / "figures"
prs = Presentation(str(SRC))
DARK, GREY, NAVY, ORANGE, GREEN, BLUE = "202020", "595959", "1F3864", "ED7D31", "548235", "2E75B6"
TABLE_STYLE = "{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}"


# ---------- helpers (same style as fill_ppt.py) --------------------------------
def is_bar(sh):
    return sh.width == prs.slide_width and (sh.top < Emu(200000) or sh.top > Emu(6600000))


def is_title(sh):
    return (not is_bar(sh) and sh.has_text_frame and sh.text_frame.text.strip() != ""
            and sh.top < Emu(500000) and sh.width > Emu(4000000))


def set_text(shape, text):
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    p0.runs[0].text = text


def textbox(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, space_after=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    bp = tf._txBody.find(qn("a:bodyPr"))
    for k in ("lIns", "tIns", "rIns", "bIns"):
        bp.set(k, "0")
    if bp.find(qn("a:spAutoFit")) is None:  # new text boxes already carry one; a second breaks the file
        etree.SubElement(bp, qn("a:spAutoFit"))
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space_after:
            p.space_after = Pt(space_after)
        for text, size, bold, color, italic in (para if isinstance(para, list) else [para]):
            r = p.add_run()
            r.text = text
            f = r.font
            f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, "Calibri"
            f.color.rgb = RGBColor.from_string(color)
    return tb


def heading(slide, x, y, w, text, size=15, color=DARK):
    return textbox(slide, x, y, w, 0.35, [(text, size, True, color, False)])


def body(slide, x, y, w, h, lines, size=12, color=DARK, space_after=3):
    paras = []
    for ln in lines:
        if isinstance(ln, list):
            paras.append([(t, size, b, color, False) for t, b in ln])
        else:
            paras.append((ln, size, False, color, False))
    return textbox(slide, x, y, w, h, paras, space_after=space_after)


def note(slide, x, y, w, text, size=11, align=PP_ALIGN.LEFT):
    return textbox(slide, x, y, w, 0.3, [(text, size, False, GREY, True)], align=align)


def accent(slide, x, y, h, color, w=0.07):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(color)
    sh.line.fill.background()
    return sh


def card(slide, x, y, w, h, title, text, color, size=11.5):
    """Coloured left bar + bold title + body text (v3 'Problem Statement' style)."""
    accent(slide, x, y, h, color)
    heading(slide, x + 0.18, y + 0.05, w - 0.2, title, size=13, color=color)
    lines = text if isinstance(text, list) else [text]
    body(slide, x + 0.18, y + 0.42, w - 0.2, h - 0.45, lines, size=size)


def table(slide, x, y, w, rows, col_w, row_h=0.27, size=11, hl_rows=(), left_cols=(0,)):
    gf = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                Inches(row_h * len(rows)))
    tbl = gf.table
    tbl._tbl.tblPr.find(qn("a:tableStyleId")).text = TABLE_STYLE
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.margin_left = c.margin_right = Inches(0.06)
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = RGBColor.from_string(NAVY)
            elif i in hl_rows:
                c.fill.fore_color.rgb = RGBColor.from_string("FCE4D6")
            else:
                c.fill.fore_color.rgb = RGBColor.from_string("F2F2F2" if i % 2 else "FFFFFF")
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j in left_cols else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(val)
            f = r.font
            f.size, f.name = Pt(size), "Calibri"
            f.bold = i == 0 or j == 0 or i in hl_rows
            f.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else DARK)
    return gf


def picture(slide, name, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(str(FIG / name), Inches(x), Inches(y), **kw)


def caption(slide, x, y, w, text):
    return note(slide, x, y, w, text, size=10, align=PP_ALIGN.CENTER)


TEMPLATE = prs.slides[10]  # v3 slide with bars + title box
BLANK = TEMPLATE.slide_layout
OLD = list(prs.slides._sldIdLst)[1:]  # v3 slides 2-18, removed at the end


def new_slide(title):
    s = prs.slides.add_slide(BLANK)
    for sh in TEMPLATE.shapes:
        if is_bar(sh) or is_title(sh):
            s.shapes._spTree.append(copy.deepcopy(sh._element))
    for sh in s.shapes:
        if is_title(sh):
            set_text(sh, title)
    return s


# =================== 1. Introduction (report ch. Introduction) ===================
s = new_slide("Introduction")
heading(s, 0.55, 1.2, 8.9, "Background and Motivation")
body(s, 0.55, 1.62, 8.9, 1.5, [
    "The ECG is a widely used non-invasive tool for detecting cardiac arrhythmias, and automating its "
    "interpretation is an active research area.",
    "• CNNs are used in 58.7% of 368 studies reviewed by Xiao et al. [1] and in 62% of 119 studies "
    "reviewed by Kulkarni et al. [2].",
    "• Hybrid CNN-recurrent and Transformer-based models are increasingly common [1], [2].",
], size=13)
heading(s, 0.55, 3.35, 8.9, "Significance of the Problem: three issues limit clinical usefulness")
card(s, 0.55, 3.9, 2.85, 1.9, "Evaluation protocol",
     "F1 fell from 95.52% (intra-patient) to 83.89% (inter-patient) [1]; only 30.3% of 122 reviewed "
     "articles followed the inter-patient paradigm [4].", NAVY)
card(s, 3.6, 3.9, 2.85, 1.9, "Black-box models",
     "SHAP and saliency methods are increasingly used, but fewer than 10% of studies validated "
     "explanation maps against clinically meaningful ECG structures [2].", ORANGE)
card(s, 6.65, 3.9, 2.85, 1.9, "Narrow robustness tests",
     "Robustness to realistic noise is tested narrowly, often with additive white Gaussian noise "
     "only [5].", GREEN)

# =================== 2. Problem statement & objectives ===================
s = new_slide("Problem Statement & Objectives")
accent(s, 0.55, 1.2, 1.05, ORANGE)
body(s, 0.75, 1.25, 8.7, 1.0, [
    "There is a need for an ECG arrhythmia classification framework whose performance is measured on "
    "unseen patients, whose explanations are checked for consistency and clinical plausibility, and whose "
    "predictions and explanations are evaluated for robustness under realistic noise."], size=14)
objs = [
    ("Patient-independent baseline",
     "Develop a hybrid deep learning arrhythmia classifier evaluated under a strict inter-patient protocol "
     "(MIT-BIH DS1/DS2) with external testing on a second database, reporting per-class and macro F1 "
     "alongside accuracy."),
    ("Validated explanations",
     "Apply post-hoc XAI methods (e.g. SHAP and Grad-CAM) to the patient-independent model and measure the "
     "within-class consistency of explanations and their agreement with recognised ECG features (P wave, "
     "QRS complex, RR interval)."),
    ("Robustness of predictions and explanations",
     "Assess how classification performance and explanation stability degrade under realistic NSTDB noise "
     "(baseline wander, muscle artefact, electrode motion) at graded SNRs, using uncertainty estimates to "
     "flag unreliable predictions."),
]
for k, (t, d) in enumerate(objs):
    y = 2.65 + k * 1.5
    textbox(s, 0.55, y, 0.7, 0.5, [(f"0{k + 1}", 26, True, ORANGE, False)])
    heading(s, 1.35, y, 8.1, f"Objective {k + 1}: {t}", size=14, color=NAVY)
    body(s, 1.35, y + 0.4, 8.1, 1.0, [d], size=12)

# =================== 3. Literature review ===================
s = new_slide("Literature Review")
note(s, 0.55, 1.12, 8.9, "A literature survey of 36 papers (5 reviews and 31 original studies) on deep learning "
                          "and explainable AI for ECG arrhythmia classification, grouped into four themes.",
     size=12)
themes = [
    ("Architectures", NAVY, [
        "• CNNs remain dominant [1], [2]; hybrids pair a CNN front end with a recurrent or attention module "
        "(CNN-LSTM [8], CNN-BiLSTM [11]).",
        "• Attention/Transformer models gain 0.845% over 1D-CNNs, at higher cost beyond ~1 M parameters [2].",
        "• Lightweight designs stay competitive: a 0.16 MB CNN-LSTM ran at 5.127 ms per rhythm on a "
        "Raspberry Pi [8]."]),
    ("Explainable AI methods", ORANGE, [
        "• SHAP is the most widely used method [8], [3]; then Grad-CAM and class activation maps.",
        "• Explanations for APB and PVC beats were inconsistent within a class and not validated by "
        "medical professionals [3].",
        "• Fewer than 10% of studies validated maps against clinical ECG structures [2]."]),
    ("Datasets and evaluation", GREEN, [
        "• MIT-BIH dominates: used in 61% [1] and 78.7% [4] of reviewed studies.",
        "• One study: 99.00% accuracy beat-wise but 91.69% patient-wise, with Fusion recall 0.021 [11].",
        "• External validation in only 18.5% of studies, with a mean accuracy drop of 8.7% [2]."]),
    ("Robustness", BLUE, [
        "• Often only AWGN tested: F1 fell to 32.46% at 5 dB [5].",
        "• Uncertainty-aware fusion was more robust under NSTDB noise from 15 to 0 dB [30].",
        "• Rejecting high-uncertainty predictions raised macro F1 from 0.6635 to 0.8688 [33]."]),
]
for k, (t, c, lines) in enumerate(themes):
    x = 0.55 if k % 2 == 0 else 5.1
    y = 1.65 if k < 2 else 4.35
    card(s, x, y, 4.35, 2.55, t, lines, c, size=11)

# =================== 4. Comparison of key studies (Table 1) ===================
s = new_slide("Comparison of Key Studies")
note(s, 0.55, 1.12, 8.9, "Surveyed papers most relevant to this project (report Table 1).", size=12)
table(s, 0.55, 1.55, 8.9, [
    ["Ref.", "Dataset(s)", "Model", "XAI", "Split", "Key result"],
    ["[8]", "MIT-BIH + LTAF", "Lightweight CNN-LSTM", "SHAP", "Not stated (85/15)",
     "Acc 98.24%, Se 86.1%; 0.16 MB"],
    ["[3]", "MIT-BIH + 12-lead (Zheng et al.)", "1D CNN", "4 SHAP variants", "Not stated (80/20)",
     "Val acc 98.30%; 73.45% on combined data"],
    ["[29]", "MIT-BIH; Chapman-Shaoxing", "Adversarial SE-ResNet", "Not stated", "Inter-patient",
     "AFib F1 77.9% vs 61.0%"],
    ["[5]", "MIT-BIH; INCART, SVDB, PTB", "CNN-BiLSTM + CapsNet fusion", "Branch weights", "Inter-patient",
     "Acc 99.55%; F1 32.46% at 5 dB AWGN"],
    ["[11]", "MIT-BIH", "CNN-BiLSTM + cGAN", "Not stated", "Both", "99.00% beat-wise vs 91.69% patient-wise"],
    ["[30]", "MIT-BIH; INCART; NSTDB", "BiLSTM + ViT, D-S fusion", "Not stated", "Not stated (80/20)",
     "Acc 98.8%; robust 15–0 dB noise"],
], [0.55, 1.75, 1.75, 1.15, 1.3, 2.4], row_h=0.55, size=11, left_cols=(0, 1, 2, 5))
body(s, 0.55, 5.6, 8.9, 1.2, [
    [("Pattern: ", True), ("XAI studies often do not describe patient separation in their splits [8], [3], while "
                          "strong inter-patient studies often include no XAI [29], [11].", False)],
], size=12)

# =================== 5. Research gaps ===================
s = new_slide("Research Gaps")
gaps = [
    ("Gap 1 – Explanations rarely validated", NAVY,
     "SHAP and Grad-CAM are widely applied, but fewer than 10% of studies validated maps against clinical "
     "ECG structures [2], and explanations were inconsistent within a class [3]."),
    ("Gap 2 – XAI seldom patient-independent", ORANGE,
     "Several XAI studies do not describe patient separation, although inter-patient evaluation "
     "substantially lowers performance [1], [11] and is used by only 30.3% of studies [4]."),
    ("Gap 3 – Narrow noise tests, not linked to XAI", GREEN,
     "Noise tests are limited to AWGN [5] or absent; NSTDB-based work [30] has no explanation analysis; no "
     "surveyed paper tests whether explanations stay stable under noise [33]."),
    ("Gap 4 – Weak minority classes and generalisation", BLUE,
     "Patient-wise Fusion and S recall were only 0.021 and 0.209 [11]; only 18.5% of studies validate "
     "externally, with a mean accuracy drop of 8.7% [2]."),
]
for k, (t, c, d) in enumerate(gaps):
    x = 0.55 if k % 2 == 0 else 5.1
    y = 1.4 if k < 2 else 4.0
    card(s, x, y, 4.35, 2.1, t, d, c, size=12.5)

# =================== 6. Methodology: block diagram + datasets ===================
s = new_slide("Methodology: Proposed Framework")
picture(s, "block_diagram.png", 2.05, 1.08, w=5.9)
caption(s, 0.55, 4.8, 8.9, "Block diagram of the proposed framework (report Figure 1)")
heading(s, 0.55, 5.15, 8.9, "Datasets", size=14)
body(s, 0.55, 5.5, 8.9, 1.5, [
    "• MIT-BIH Arrhythmia Database with AAMI classes N, S, V, F, Q and the inter-patient DS1/DS2 split "
    "(Objective 1).",
    "• External database for cross-database testing (INCART preferred: beat-level annotations like "
    "MIT-BIH); a stretch goal, otherwise future work.",
    "• MIT-BIH Noise Stress Test Database (baseline wander, muscle artefact, electrode motion) for "
    "robustness tests (Objective 3), as used in [30].",
], size=11.5)

# =================== 7. Classification model & training ===================
s = new_slide("Classification Model & Training")
note(s, 0.55, 1.12, 8.9, "A hybrid model: convolutional front end + recurrent or attention module for temporal "
                          "context, a design family a recent review judged the most clinically translatable [2].",
     size=12)
mods = [
    ("Baseline: 1D CNN", NAVY,
     "Three conv blocks (32, 64, 128 filters; kernels 7, 5, 3; each with batch normalisation, ReLU and "
     "max-pooling), global average pooling, Dense 64 and softmax (about 44,700 parameters). Already "
     "implemented and trained."),
    ("Proposed: CNN-LSTM + RR", ORANGE,
     "Same conv front end → LSTM (64 units); its last hidden state is concatenated with the previous and next "
     "RR intervals (each divided by the record's mean RR) → Dense 64 → softmax over N, S, V, F. A residual "
     "CNN is kept as an extra comparison model."),
    ("Class imbalance", GREEN,
     "Class-weighted cross-entropy on the training partition only. Balanced weights (F:N 113:1) caused many "
     "false Fusion predictions; square-root weights (10.6:1) gave the best results and are used."),
    ("Training", BLUE,
     "Adam, batch size 128, at most 30 epochs; best epoch chosen on validation macro F1 rather than "
     "validation loss (this alone raised test macro F1 from 76.29% to 83.96%). Learning rate tuned on "
     "validation."),
]
for k, (t, c, d) in enumerate(mods):
    x = 0.55 if k % 2 == 0 else 5.1
    y = 1.75 if k < 2 else 3.95
    card(s, x, y, 4.35, 1.95, t, d, c, size=12)
body(s, 0.55, 6.35, 8.9, 0.6, [
    [("Tools: ", True), ("Python 3, TensorFlow 2.20 / Keras, NumPy, pandas, scikit-learn, Matplotlib, seaborn, "
                        "wfdb, SHAP, Streamlit; Kaggle notebooks with an NVIDIA Tesla T4 GPU.", False)]], size=11)

# =================== 8. Explainability, robustness & metrics ===================
s = new_slide("Explainability, Robustness & Metrics")
body(s, 0.55, 1.15, 8.9, 0.7, [
    [("Explainability: ", True), ("Grad-CAM [37] (already implemented) and Gradient SHAP [38], evaluated for "
                                 "within-class consistency and agreement with the P wave, QRS complex and RR "
                                 "interval. Three measures:", False)]], size=12.5)
xs = [
    ("Region share", NAVY,
     "Share of total absolute attribution in fixed windows around the R peak: P [−250, −80] ms, QRS "
     "[−50, +50] ms, T [+100, +400] ms; averaged per class."),
    ("Deletion test", ORANGE,
     "Zero the 10% highest-attribution samples; compare the drop in predicted-class probability with "
     "removing 10% at random. A faithful map gives a much larger drop."),
    ("Stability", GREEN,
     "Pearson correlation between the attribution map of a clean beat and of the same beat with noise "
     "added (used in the robustness study)."),
]
for k, (t, c, d) in enumerate(xs):
    card(s, 0.55 + k * 3.05, 2.05, 2.85, 1.95, t, d, c, size=11.5)
heading(s, 0.55, 4.45, 4.35, "Robustness and uncertainty", size=14)
body(s, 0.55, 4.85, 4.35, 2.0, [
    "• NSTDB muscle-artefact ('ma') record scaled and added to test beats at SNR = 18, 6 and 0 dB, "
    "SNR = 10·log10(P_signal / P_noise).",
    "• Macro F1 plotted against SNR for both models.",
    "• Monte Carlo dropout [33] to flag unreliable predictions if time permits; otherwise future work.",
], size=11.5)
heading(s, 5.1, 4.45, 4.35, "Evaluation metrics", size=14)
body(s, 5.1, 4.85, 4.35, 2.0, [
    "• Accuracy, per-class precision, recall and F1, and macro F1.",
    "• Overall accuracy alone can hide poor minority-class performance [1], so macro F1 is the main metric.",
], size=11.5)

# =================== 9. Work done till mid-semester ===================
s = new_slide("Work Done Till Mid-Semester")
heading(s, 0.55, 1.15, 4.3, "Completed Tasks")
body(s, 0.55, 1.55, 4.3, 2.5, [
    "• Literature survey and identification of the research gaps.",
    "• Framework design: AAMI class mapping, intra- and inter-patient (DS1/DS2) protocols, explanation "
    "metrics and noise test settings.",
    "• V1: baseline 1D CNN with class-weighted loss, full test-set evaluation and Grad-CAM for each class.",
    [("• V2: five controlled experiments (E1–E5) + threshold tuning: test macro F1 ", False),
     ("76.29% → 91.10%", True), (", then Grad-CAM on the final model.", False)],
], size=11.5)
heading(s, 5.1, 1.15, 4.35, "Dataset used in the preliminary experiments")
body(s, 5.1, 1.55, 4.35, 2.5, [
    "Pre-segmented MIT-BIH heartbeat dataset of Kachuee et al. [35]: 187-sample beats (1.5 s at 125 Hz), "
    "scaled to [0, 1], zero-padded, five AAMI classes (N, S, V, F, Q).",
    "Used to develop and debug the pipeline before building our own pre-processing.",
], size=11.5)
table(s, 0.55, 4.15, 8.9, [
    ["Class", "Training file", "%", "Test file", "%"],
    ["N – Normal", "72,471", "82.77", "18,118", "82.76"],
    ["S – Supraventricular ectopic", "2,223", "2.54", "556", "2.54"],
    ["V – Ventricular ectopic", "5,788", "6.61", "1,448", "6.61"],
    ["F – Fusion", "641", "0.73", "162", "0.74"],
    ["Q – Unknown / paced", "6,431", "7.35", "1,608", "7.35"],
    ["Total", "87,554", "100", "21,892", "100"],
], [3.3, 1.6, 1.0, 1.6, 1.4], row_h=0.27, size=11, hl_rows=(6,))
note(s, 0.55, 6.15, 8.9, "A model that always predicts N would already exceed 80% accuracy, so macro F1 is the "
                          "main metric. The split is by beat, so all preliminary results are intra-patient.",
     size=11)

# =================== 10. Preliminary data: split + figures ===================
s = new_slide("Preliminary Data: Split and Example Beats")
table(s, 0.55, 1.12, 8.9, [
    ["Set", "Beats", "Input shape", "Source", "Used for"],
    ["Training", "70,043", "(187, 1)", "80% of training file (stratified, seed 42)", "Fitting the weights"],
    ["Validation", "17,511", "(187, 1)", "20% of training file", "Epoch, model and threshold selection"],
    ["Test", "21,892", "(187, 1)", "Provided test file", "Final evaluation only"],
], [1.2, 0.9, 1.1, 3.0, 2.7], row_h=0.28, size=11)
picture(s, "V1_class_distribution.png", 0.55, 2.4, w=3.9)
caption(s, 0.55, 4.8, 3.9, "Class distribution of the training file")
body(s, 4.75, 2.55, 4.7, 2.3, [
    "• Training portion after the split: 57,977 N, 1,778 S, 4,630 V, 513 F and 5,145 Q beats.",
    "• Every beat starts at its R peak, contains a second large peak (most likely the next R peak) and is "
    "then zero-padded to 187 samples.",
    "• The notebooks perform no filtering or segmentation of their own; this was done by the dataset "
    "authors [35].",
], size=11.5)
picture(s, "V1_example_beats.png", 0.55, 5.1, w=8.9)
caption(s, 0.55, 6.6, 8.9, "One example beat of each class (187 samples, 125 Hz, zero-padded)")

# =================== 11. V1 architecture & settings ===================
s = new_slide("V1: Baseline 1D CNN – Architecture & Training")
table(s, 0.55, 1.12, 8.9, [
    ["Layer", "Output shape", "Parameters"],
    ["Input (one beat)", "(187, 1)", "0"],
    ["Conv1D 32, k = 7 → BN → ReLU → MaxPool 2", "(93, 32)", "256 + 128"],
    ["Conv1D 64, k = 5 → BN → ReLU → MaxPool 2", "(46, 64)", "10,304 + 256"],
    ["Conv1D 128, k = 3 → BN → ReLU (Grad-CAM layer) → MaxPool 2", "(23, 128)", "24,704 + 512"],
    ["Global average pooling", "(128)", "0"],
    ["Dense 64, ReLU → Dropout 0.3", "(64)", "8,256"],
    ["Dense 5, softmax", "(5)", "325"],
    ["Total (trainable / non-trainable)", "", "44,741 (44,293 / 448)"],
], [5.3, 1.6, 2.0], row_h=0.27, size=11, hl_rows=(8,))
table(s, 0.55, 3.75, 8.9, [
    ["Setting", "Value"],
    ["Optimiser / learning rate", "Adam, 10⁻³"],
    ["Loss", "Sparse categorical cross-entropy with balanced class weights"],
    ["Class weights (N, S, V, F, Q)", "0.24, 7.88, 3.03, 27.31, 2.72 (F:N ≈ 113:1)"],
    ["Batch size / maximum epochs", "128 / 30"],
    ["Early stopping", "Validation loss, patience 6, best weights restored"],
    ["Learning-rate schedule", "ReduceLROnPlateau on validation loss, factor 0.5, patience 3, min 10⁻⁵"],
    ["Augmentation / random seed", "None / 42 (Python, NumPy, TensorFlow)"],
], [2.9, 6.0], row_h=0.27, size=11)
note(s, 0.55, 6.05, 8.9, "Class weights w_c = N / (K·n_c). Training stopped after 17 epochs and the weights of "
                          "epoch 11 were restored.", size=11)

# =================== 12. V1 results ===================
s = new_slide("V1 Result: Baseline 1D CNN")
picture(s, "V1_confusion_matrix.png", 0.9, 1.1, w=8.2)
table(s, 0.55, 4.45, 4.4, [
    ["Class", "Precision", "Recall", "F1", "Beats"],
    ["N", "0.9813", "0.9594", "0.9702", "18,118"],
    ["S", "0.5586", "0.6691", "0.6088", "556"],
    ["V", "0.8713", "0.8605", "0.8659", "1,448"],
    ["F", "0.2602", "0.8642", "0.4000", "162"],
    ["Q", "0.9896", "0.9502", "0.9695", "1,608"],
    ["Macro", "0.7322", "0.8607", "0.7629", "21,892"],
], [0.8, 0.9, 0.9, 0.9, 0.9], row_h=0.27, size=11, hl_rows=(6,))
heading(s, 5.15, 4.4, 4.3, "Accuracy 94.41%  |  Macro F1 76.29%", size=14)
body(s, 5.15, 4.82, 4.3, 2.0, [
    "• Class weighting gave high rare-class recall (F 86.4%) at the cost of precision.",
    "• Of 538 beats predicted F, only 140 were truly F (293 were N): F precision 26.0%.",
    "• Largest S error: S → N (163 of 556 S beats, 29%).",
    "• Validation accuracy jumped between 14.5% and 95.1% from epoch to epoch.",
], size=11)

# =================== 13. V1 training behaviour ===================
s = new_slide("V1: Training Curves and Per-Class Scores")
picture(s, "V1_training_curves.png", 0.55, 1.1, w=8.9)
caption(s, 0.55, 4.0, 8.9, "Training and validation accuracy and loss (17 epochs; best validation loss at epoch 11)")
picture(s, "V1_per_class_scores.png", 0.55, 4.35, w=4.3)
body(s, 5.1, 4.45, 4.35, 2.5, [
    "• Training loss fell smoothly 0.706 → 0.175; training accuracy still rising (90.93% at epoch 17).",
    "• Validation loss ranged from 0.349 (epoch 11, restored) to 4.29 (epoch 2): a gap between training and "
    "inference behaviour, not overfitting – what V2 set out to fix.",
    "• F and S precision bars are clearly below recall: the sign of over-weighting the rare classes.",
    "• Weighted F1 94.98%; macro precision / recall 73.22% / 86.07%.",
], size=11)

# =================== 14. V2 setup ===================
s = new_slide("V2: Step-by-Step Experiment Setup")
table(s, 0.55, 1.1, 4.5, [
    ["Class", "Train beats", "Balanced (B0–E2)", "Sqrt (E3–E5)"],
    ["N", "57,977", "0.24", "0.49"],
    ["S", "1,778", "7.88", "2.81"],
    ["V", "4,630", "3.03", "1.74"],
    ["F", "513", "27.31", "5.23"],
    ["Q", "5,145", "2.72", "1.65"],
    ["F : N", "", "113 : 1", "10.6 : 1"],
], [0.75, 1.05, 1.4, 1.3], row_h=0.25, size=10.5, hl_rows=(6,))
body(s, 5.3, 1.1, 4.15, 1.9, [
    "Same split and seed as V1; one change added per step. The test set was never used to choose an "
    "epoch, a model or a threshold.",
    "Custom loop: best weights stored on validation macro F1; LR halved after 3 epochs without improvement "
    "(min 10⁻⁵); stop after 8.",
], size=11)
table(s, 0.55, 3.0, 8.9, [
    ["ID", "Network", "Params", "LR", "Clip norm", "BN mom.", "Class weights", "Augment."],
    ["E1", "CNN", "44,741", "1×10⁻³", "–", "0.99", "Balanced", "No"],
    ["E2", "CNN", "44,741", "5×10⁻⁴", "1.0", "0.9", "Balanced", "No"],
    ["E3", "CNN", "44,741", "5×10⁻⁴", "1.0", "0.9", "Square-root", "No"],
    ["E4", "CNN", "44,741", "5×10⁻⁴", "1.0", "0.9", "Square-root", "Yes"],
    ["E5", "Residual CNN [36]", "185,477", "5×10⁻⁴", "1.0", "0.9", "Square-root", "Yes"],
], [0.6, 1.75, 1.0, 1.0, 1.0, 0.95, 1.4, 1.2], row_h=0.25, size=10.5)
picture(s, "V2_01_augmentation_examples.png", 0.55, 4.6, w=5.6)
caption(s, 0.55, 5.7, 5.6, "Original and three augmented copies of an N, S and F beat")
body(s, 6.35, 4.6, 3.1, 2.4, [
    [("Augmentation (E4, E5): ", True), ("amplitude ×U(0.9, 1.1), Gaussian noise σ = 0.01 on the non-padded "
                                        "part, shift ±3 samples (±24 ms).", False)],
    [("Residual CNN: ", True), ("Conv 32 (k 7) stem + 3 residual blocks (32, 64, 128), two Conv1D (k 5) + BN "
                               "per block with a shortcut.", False)],
], size=10.5)

# =================== 15. V2 training runs ===================
s = new_slide("V2: Training Runs")
table(s, 0.55, 1.1, 8.9, [
    ["ID", "Best epoch / epochs run", "Stopped early", "Best val. macro F1 (%)", "Δ test macro F1", "Time (s)"],
    ["E1", "16 / 24", "Yes", "85.43", "+7.67 (vs B0)", "70"],
    ["E2", "28 / 30", "No (cap)", "79.35", "−5.66", "80"],
    ["E3", "23 / 30", "No (cap)", "88.65", "+9.28", "80"],
    ["E4", "25 / 30", "No (cap)", "88.01", "−1.38", "95"],
    ["E5", "18 / 26", "Yes", "93.19", "+4.57", "158"],
], [0.7, 2.0, 1.4, 2.0, 1.6, 1.2], row_h=0.27, size=11, hl_rows=(5,))
picture(s, "V2_02_validation_curves.png", 0.55, 2.85, w=8.9)
caption(s, 0.55, 5.35, 8.9, "Validation macro F1 and validation accuracy per epoch for E1–E5")
body(s, 0.55, 5.75, 8.9, 1.2, [
    "• E5 trains highest and most smoothly (val. macro F1 0.78 after epoch 1; val. accuracy 0.94–0.99); "
    "E1 is the most jagged and E2 the lowest.",
    "• E2, E3 and E4 reached the 30-epoch cap, so they may be slightly under-trained. All five runs took "
    "8.1 min on a Tesla T4 GPU.",
], size=11)

# =================== 16. V2 results table ===================
s = new_slide("V2: Step-by-Step Results")
table(s, 0.55, 1.1, 8.9, [
    ["ID", "Change", "Acc.", "Macro F1", "S F1", "F F1", "Prec. F", "Rec. F"],
    ["B0", "V1 baseline", "94.41", "76.29", "60.90", "40.00", "26.02", "86.42"],
    ["E1", "Epoch chosen on val. macro F1", "96.28", "83.96", "63.48", "68.12", "55.95", "87.04"],
    ["E2", "+ lower LR, clipping, BN 0.9", "93.58", "78.30", "49.13", "58.49", "43.73", "88.27"],
    ["E3", "+ square-root class weights", "97.39", "87.58", "69.84", "76.97", "75.60", "78.40"],
    ["E4", "+ light augmentation", "97.02", "86.20", "66.16", "76.25", "72.63", "80.25"],
    ["E5", "+ residual network", "98.29", "90.77", "81.22", "78.53", "78.05", "79.01"],
    ["Final", "E5 + tuned thresholds", "98.41", "91.10", "82.14", "79.23", "82.12", "76.54"],
], [0.7, 2.8, 0.8, 1.0, 0.85, 0.85, 0.95, 0.95], row_h=0.3, size=11, hl_rows=(7,), left_cols=(0, 1))
note(s, 0.55, 3.55, 8.9, "Test set, %. Final = E5 with per-class probability scale factors tuned on validation "
                          "(only N × 3 changed).", size=11)
heading(s, 0.55, 4.0, 4.35, "What helped", color=GREEN)
body(s, 0.55, 4.4, 4.35, 2.5, [
    "• E1: selecting the epoch on validation macro F1: +7.67 macro F1.",
    "• E3: square-root class weights: +9.28 over E2 (largest gain).",
    "• E5: residual network: +4.57 over E4.",
], size=12)
heading(s, 5.1, 4.0, 4.35, "What did not", color=ORANGE)
body(s, 5.1, 4.4, 4.35, 2.5, [
    "• E2 stability settings: −5.66; validation swings not visibly removed.",
    "• E4 augmentation: −1.38.",
    "• Threshold tuning: only +0.33; traded some S and F recall for precision.",
], size=12)

# =================== 17. V2 per-class progression ===================
s = new_slide("V2: Per-Class Progression")
picture(s, "V2_03_experiment_comparison.png", 0.55, 1.1, w=4.35)
caption(s, 0.55, 2.95, 4.35, "Test macro F1, S recall and F recall per step")
picture(s, "V2_08_per_class_f1_progression.png", 5.1, 1.1, w=4.35)
caption(s, 5.1, 3.45, 4.35, "Per-class F1 across the V2 experiments")
table(s, 0.55, 3.85, 8.9, [
    ["Model", "Acc.", "Macro F1", "N", "S", "V", "F", "Q"],
    ["B0", "94.41", "76.29", "97.00", "60.90", "86.60", "40.00", "97.00"],
    ["E1", "96.28", "83.96", "97.92", "63.48", "92.51", "68.12", "97.77"],
    ["E2", "93.58", "78.30", "96.27", "49.13", "90.77", "58.49", "96.84"],
    ["E3", "97.39", "87.58", "98.58", "69.84", "94.19", "76.97", "98.29"],
    ["E4", "97.02", "86.20", "98.40", "66.16", "91.93", "76.25", "98.29"],
    ["E5", "98.29", "90.77", "99.13", "81.22", "95.70", "78.53", "99.28"],
    ["Final", "98.41", "91.10", "99.18", "82.14", "95.87", "79.23", "99.06"],
], [1.1, 1.1, 1.2, 1.1, 1.1, 1.1, 1.1, 1.1], row_h=0.3, size=11, hl_rows=(7,))
note(s, 0.55, 6.35, 8.9, "Per-class test F1 (%) of every model.", size=10.5)

# =================== 18. Final model ===================
s = new_slide("Final Model: Residual CNN + Tuned Thresholds")
picture(s, "V2_04_confusion_matrix_final.png", 0.9, 1.1, w=8.2)
table(s, 0.55, 4.45, 4.4, [
    ["Class", "Precision", "Recall", "F1", "Beats"],
    ["N", "0.9874", "0.9962", "0.9918", "18,118"],
    ["S", "0.8758", "0.7734", "0.8214", "556"],
    ["V", "0.9805", "0.9378", "0.9587", "1,448"],
    ["F", "0.8212", "0.7654", "0.7923", "162"],
    ["Q", "0.9981", "0.9832", "0.9906", "1,608"],
    ["Macro", "0.9326", "0.8912", "0.9110", "21,892"],
], [0.8, 0.9, 0.9, 0.9, 0.9], row_h=0.27, size=11, hl_rows=(6,))
heading(s, 5.15, 4.4, 4.3, "Accuracy 98.41%  |  Macro F1 91.10%", size=14)
body(s, 5.15, 4.82, 4.3, 2.0, [
    "• Selected model: E5, highest validation macro F1 (93.19%).",
    "• Tuned scale factors: N = 3.0, all others 1.0, so non-N predictions need more confidence.",
    "• Tuning raised accuracy 98.29% → 98.41% and macro F1 90.77% → 91.10%.",
    "• Weighted-average F1: 98.37%.",
], size=11)

# =================== 19. Baseline vs final ===================
s = new_slide("Baseline (V1) vs Final Model (V2)")
picture(s, "V2_09_final_per_class_metrics.png", 0.55, 1.1, w=4.2)
caption(s, 0.55, 3.62, 4.2, "Per-class P / R / F1 of the final model")
table(s, 4.95, 1.1, 4.5, [
    ["Metric", "B0", "Final", "Change"],
    ["Accuracy (%)", "94.41", "98.41", "+4.00"],
    ["Macro F1 (%)", "76.29", "91.10", "+14.81"],
    ["Weighted F1 (%)", "94.98", "98.37", "+3.39"],
    ["Precision / Recall S", "55.86 / 66.91", "87.58 / 77.34", "+31.72 / +10.43"],
    ["Precision / Recall F", "26.02 / 86.42", "82.12 / 76.54", "+56.10 / −9.88"],
    ["Wrongly predicted F", "398", "27", "−371"],
    ["Wrongly predicted S", "294", "61", "−233"],
    ["S predicted as N", "163 (29%)", "120 (22%)", "−43"],
    ["Parameters", "44,741", "185,477", "≈ 4×"],
], [1.5, 0.95, 0.95, 1.1], row_h=0.27, size=9.5, hl_rows=(2,))
body(s, 0.55, 4.15, 8.9, 2.7, [
    "• F1 improved for every class: S 60.9 → 82.1, V 86.6 → 95.9, F 40.0 → 79.2.",
    "• Fusion recall is lower than the baseline's (76.5% vs 86.4%), because the baseline's high recall came "
    "from over-predicting F.",
    "• Largest remaining error: S → N (120 of 556 S beats, 22%). S beats differ from N mainly in timing, "
    "which supports adding explicit RR-interval features in the proposed model.",
], size=12)

# =================== 20. Grad-CAM figures ===================
s = new_slide("Grad-CAM Explanations")
heading(s, 0.55, 1.08, 8.9, "V1 baseline CNN", size=13)
picture(s, "V1_gradcam.png", 0.55, 1.4, w=8.9)
heading(s, 0.55, 3.3, 8.9, "V2 final model (residual CNN)", size=13)
picture(s, "V2_06_gradcam_final.png", 0.55, 3.6, w=8.9)
body(s, 0.55, 5.55, 8.9, 1.4, [
    "• Grad-CAM [37] on the last convolutional activation (46 time steps, interpolated to 187 samples); the "
    "most confidently and correctly classified test beat of each class.",
    "• Baseline: mostly the window start and the second large peak / end of signal. Final model: broader "
    "maps that more often cover bumps and troughs.",
], size=11)

# =================== 21. Grad-CAM table ===================
s = new_slide("Grad-CAM: Where the Heat Is Concentrated")
table(s, 0.55, 1.1, 8.9, [
    ["Class", "V1 baseline CNN", "Final model (residual CNN)"],
    ["N", "Window start (0–0.1 s) and 1.15–1.3 s around the second peak / start of padding",
     "0.12–0.2 s and 0.4–0.5 s, before the second R peak (P-wave-like bump); second peak itself low"],
    ["S", "0.35–0.45 s over the second (sharp) peak and the drop into padding",
     "Window start (0–0.2 s) and 0.75–0.9 s just after the second peak"],
    ["V", "Window start, broad 0.3–0.9 s region, strong band at 1.0–1.1 s where the signal ends",
     "Window start, 0.27–0.35 s, and 1.15–1.3 s around the second peak / end of signal"],
    ["F", "Window start and two narrow bands either side of the second peak (~0.4 and ~0.53 s)",
     "One strong broad band at 0.3–0.5 s centred on the second peak"],
    ["Q", "Window start and a strong band at 0.75–0.85 s over the second peak",
     "~0.1 s (deep trough) and 0.75–0.95 s around the second peak"],
], [0.7, 4.1, 4.1], row_h=0.55, size=11, left_cols=(0, 1, 2))
note(s, 0.55, 4.5, 8.9, "Read from the heatmap figures; all beats predicted with p = 1.00.", size=10.5)
body(s, 0.55, 4.9, 8.9, 2.0, [
    "• Each window starts at the current R peak, so the second peak is likely the next R peak: the models "
    "may partly use the timing of the next beat (an implicit RR cue) as well as beat shape.",
    "• One Grad-CAM step ≈ 4 samples (≈ 32 ms); heat at the very first samples may partly be an edge effect "
    "of 'same' padding.",
    "• Qualitative only: to be tested with the region-share and deletion measures.",
], size=11.5)

# =================== 22. Challenges ===================
s = new_slide("Challenges Faced")
ch = [
    ("Severe class imbalance", NAVY,
     "Fully balanced class weights made the baseline over-predict the rare classes (F precision 26.0%). "
     "Softer, square-root weights fixed most of this."),
    ("Unstable training", ORANGE,
     "The baseline's validation accuracy swung between 14.5% and 95.1%, so early stopping on validation loss "
     "stopped training too early. Choosing the epoch on validation macro F1 and the residual network gave "
     "smoother, better results."),
    ("Not every improvement helped", GREEN,
     "The stability settings (E2) and augmentation (E4) lowered macro F1; because the experiments were "
     "cumulative, the best combination of changes is not yet known."),
    ("Optimistic preliminary setting", BLUE,
     "Pre-processed, intra-patient data including the Q class, so results are optimistic and not directly "
     "comparable with the planned four-class, inter-patient results. Each experiment was run once with one "
     "seed."),
]
for k, (t, c, d) in enumerate(ch):
    x = 0.55 if k % 2 == 0 else 5.1
    y = 1.4 if k < 2 else 4.0
    card(s, x, y, 4.35, 2.1, t, d, c, size=12.5)

# =================== 23. Work plan ===================
s = new_slide("Work Plan for Next Phase")
note(s, 0.55, 1.15, 8.9, "The following tasks follow from the project objectives and build on the pipeline "
                          "developed so far.", size=12.5)
steps = [
    "Port the pipeline from the pre-segmented Kaggle beats to the raw MIT-BIH dataset.",
    "Implement pre-processing and the inter-patient DS1/DS2 split of the MIT-BIH database.",
    "Implement and train the hybrid classifier; report per-class and macro F1; test on the external database.",
    "Apply XAI methods and evaluate explanation consistency.",
    "Generate NSTDB noise at graded SNRs; evaluate classification, explanation stability and uncertainty.",
    "Analyse results, prepare the final report and presentation.",
]
for k, t in enumerate(steps):
    y = 1.75 + k * 0.82
    textbox(s, 0.75, y, 0.6, 0.5, [(f"{k + 1}", 22, True, ORANGE, False)])
    accent(s, 1.3, y + 0.02, 0.5, NAVY, w=0.05)
    body(s, 1.5, y + 0.08, 7.9, 0.5, [t], size=14)

# =================== 24. Expected outcomes ===================
s = new_slide("Expected Outcomes")
heading(s, 0.55, 1.15, 8.9, "Deliverables")
body(s, 0.55, 1.55, 8.9, 1.8, [
    "• A patient-independent ECG arrhythmia classifier with per-class and macro-averaged results under an "
    "inter-patient protocol and on an external database.",
    "• A quantitative evaluation of explanation with two XAI methods.",
    "• A robustness and uncertainty analysis of both predictions and explanations under realistic NSTDB noise.",
    "• Final report and presentation.",
], size=13)
heading(s, 0.55, 3.55, 8.9, "Expected Results")
body(s, 0.55, 3.95, 8.9, 2.9, [
    "• Intra-patient: macro F1 of about 90%, in line with the 91.10% already reached.",
    "• Inter-patient (DS1/DS2): clearly lower, as reported in the literature [1], [11]; S and F remain the "
    "hardest classes; RR features expected to improve S recall in particular.",
    "• Explanations: V beats should place a larger share of importance on the QRS region than N beats; "
    "deleting the most important samples should lower confidence much more than deleting random ones.",
    "• Under NSTDB noise, macro F1 and explanation stability expected to fall as SNR decreases from 18 dB "
    "to 0 dB.",
], size=13)

# =================== 25. Summary (from the abstract) ===================
s = new_slide("Summary")
body(s, 0.55, 1.25, 8.9, 5.5, [
    "• Reported accuracies above 98% fall considerably on unseen patients, external databases or noisy "
    "signals, and explanations are rarely validated against clinical ECG knowledge [1], [2], [3].",
    "• A working training, evaluation and Grad-CAM pipeline has been built and tested on the pre-segmented "
    "MIT-BIH dataset [35] (five AAMI classes, intra-patient split).",
    [("• Baseline 1D CNN: ", False), ("94.41% accuracy, 76.29% macro F1", True),
     (" – very low precision on Fusion (26.0%) and S (55.9%).", False)],
    [("• Epoch selection on validation macro F1, square-root class weights and a residual network gave the "
      "largest gains: ", False), ("98.41% accuracy, 91.10% macro F1", True),
     (" on the same 21,892-beat test set.", False)],
    "• Grad-CAM heatmaps were produced for every class.",
    "• These results are intra-patient and are treated as a development baseline for the "
    "patient-independent framework.",
], size=14, space_after=8)

# =================== 26-27. References (report numbering) ===================
REFS = {
    1: "Q. Xiao, K. Lee, S. A. Mokhtar, I. Ismail, A. L. bin Md Pauzi, Q. Zhang, and P. Y. Lim, \"Deep "
       "learning-based ECG arrhythmia classification: A systematic review,\" Applied Sciences, vol. 13, Art. "
       "no. 4964, 2023.",
    2: "N. N. Kulkarni, Nagaraja G. S., B. G. Sudarshan, and Yeriswamy M. C., \"Recent advances in deep "
       "learning based arrhythmia classification using ECG and PPG signals: A systematic review,\" Discover "
       "Computing, vol. 29, Art. no. 666, 2026.",
    3: "J. Beck and A. John, \"Explainable AI (XAI) for arrhythmia detection from electrocardiograms,\" "
       "arXiv:2508.17294, 2025.",
    4: "G. A. L. Silva, P. H. L. Silva, G. J. P. Moreira, V. L. S. Freitas, J. C. Gertrudes, and E. J. S. "
       "Luz, \"A systematic review of ECG arrhythmia classification: Adherence to standards, fair evaluation, "
       "and embedded feasibility,\" arXiv:2503.07276, 2025.",
    5: "B. A. Hassoon, S. Xiong, M. A. Hasson, R. Salahudeen, A. O. Abdulsalami, and T. Khan, \"Dual-branch "
       "bidirectional attention fusion of temporal and hierarchical representations for robust ECG "
       "arrhythmia classification,\" Expert Systems with Applications, vol. 331, Art. no. 133409, 2026.",
    8: "N. Alamatsaz, L. Tabatabaei, M. Yazdchi, H. Payan, N. Alamatsaz, and F. Nasimi, \"A lightweight "
       "hybrid CNN-LSTM explainable model for ECG-based arrhythmia detection,\" Biomedical Signal Processing "
       "and Control, vol. 90, Art. no. 105884, 2024.",
    11: "Sa. I. Ibrahim, \"A hybrid deep learning algorithm for ECG-based heart disease classification,\" "
        "Scientific Reports, vol. 16, Art. no. 29911, 2026.",
    29: "Y. Jeong, J. Lee, and M. Shin, \"Enhancing inter-patient performance for arrhythmia classification "
        "with adversarial learning using beat-score maps,\" Applied Sciences, vol. 14, Art. no. 7227, 2024.",
    30: "M. Ashhad, S. Rahmani, M. Fayiz, A. Etemad, and J. Hashemi, \"Uncertainty-aware multi-view "
        "arrhythmia classification from ECG,\" arXiv:2506.06342, 2025 (published at IJCNN 2024).",
    33: "W. Zhang, X. Di, G. Wei, S. Geng, Z. Fu, and S. Hong, \"Cardiac arrhythmia classification with "
        "rejection of ECG recordings based on uncertainty estimation from deep neural networks,\" Neural "
        "Computing and Applications, vol. 36, pp. 4047–4058, 2024.",
    35: "M. Kachuee, S. Fazeli, and M. Sarrafzadeh, \"ECG heartbeat classification: A deep transferable "
        "representation,\" in Proc. IEEE Int. Conf. Healthcare Informatics (ICHI), 2018, pp. 443–444.",
    36: "K. He, X. Zhang, S. Ren, and J. Sun, \"Deep residual learning for image recognition,\" in Proc. IEEE "
        "Conf. Computer Vision and Pattern Recognition (CVPR), 2016, pp. 770–778.",
    37: "R. R. Selvaraju, M. Cogswell, A. Das, R. Vedantam, D. Parikh, and D. Batra, \"Grad-CAM: Visual "
        "explanations from deep networks via gradient-based localization,\" in Proc. IEEE Int. Conf. Computer "
        "Vision (ICCV), 2017, pp. 618–626.",
    38: "S. M. Lundberg and S.-I. Lee, \"A unified approach to interpreting model predictions,\" in Proc. "
        "Advances in Neural Information Processing Systems (NeurIPS), 2017, pp. 4765–4774.",
}
keys = sorted(REFS)
for part, chunk in enumerate((keys[:7], keys[7:])):
    s = new_slide("References" if part == 0 else "References (contd.)")
    body(s, 0.55, 1.2, 8.9, 5.8, [f"[{k}] {REFS[k]}" for k in chunk], size=12, space_after=8)
    if part == 1:
        note(s, 0.55, 6.75, 8.9, "Numbering follows the reference list of the mid-semester report.", size=10)

# ---------- remove v3 slides 2-18 ---------------------------------------------
ids = prs.slides._sldIdLst
for el in OLD:
    prs.part.drop_rel(el.get(qn("r:id")))
    ids.remove(el)

# ---------- check every [n] used on slides is in the reference list -----------
import re
used = set()
for sl in list(prs.slides)[1:-2]:
    for sh in sl.shapes:
        texts = []
        if sh.has_text_frame:
            texts.append(sh.text_frame.text)
        if getattr(sh, "has_table", False) and sh.has_table:
            texts += [c.text for r in sh.table.rows for c in r.cells]
        for t in texts:
            used |= {int(n) for n in re.findall(r"\[(\d+)\]", t)}
missing = used - set(REFS)
assert not missing, f"cited but not listed: {missing}"
print("unused refs:", set(REFS) - used)

prs.save(str(OUT))
print("saved", OUT.name, "slides:", len(prs.slides))
